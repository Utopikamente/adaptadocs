"""Genera el Anexo III.b — Adaptación Curricular Individualizada y Significativa
(ACIS) de la Comunidad de Madrid (Educación Secundaria Obligatoria).

Replica la estructura del modelo oficial y la rellena con la información que
aporta el profesorado. La herramienta **propone**; el equipo docente decide:
todo campo vacío se marca como pendiente y el documento sale como BORRADOR.

Límite legal (art. 10 del Decreto 23/2023; art. 10.3.a de la Orden 1712/2023):
la ACS está reservada al alumnado con necesidades educativas especiales, lo
determina el equipo de orientación mediante evaluación psicopedagógica. Por
eso cada materia lleva una casilla `acs_determinada`: si no está marcada, ese
bloque no se genera y se devuelve un aviso.

Los textos fijos están en `acis.json`, editable sin tocar el código.
"""

from __future__ import annotations

import json
import os
from dataclasses import dataclass, field
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

_RUTA_DATOS = os.path.join(os.path.dirname(__file__), "acis.json")


def cargar_textos() -> dict:
    """Lee `acis.json` (textos fijos y ayudas del formulario, editables)."""
    with open(_RUTA_DATOS, encoding="utf-8") as f:
        return json.load(f)


# --------------------------------------------------------------------------- #
# Modelo de datos
# --------------------------------------------------------------------------- #

@dataclass
class MateriaACIS:
    """Bloque de ACIS de una materia. Los campos de texto vacíos se marcan
    como pendientes en el documento."""

    materia: str = ""
    profesor: str = ""
    departamento: str = ""
    # Casilla imprescindible: el equipo de orientación ha determinado ACS
    # para esta materia mediante evaluación psicopedagógica.
    acs_determinada: bool = False
    competencias: str = ""            # adaptación de competencias (opcional)
    criterios_evaluacion: str = ""
    contenidos: str = ""
    metodologia: str = ""
    instrumentos: str = ""
    unidades: str = ""
    # Lista de (unidad didáctica, trimestre)
    secuenciacion: list[tuple[str, str]] = field(default_factory=list)


@dataclass
class DatosACIS:
    localidad: str = ""
    fecha: str = ""                   # p. ej. "15 de octubre de 2026"
    profesor_de: str = ""             # "EL PROFESOR DE ___"
    departamento_vb: str = ""         # "DE ___" (jefatura de departamento)
    materias: list[MateriaACIS] = field(default_factory=list)


def materia_desde_programacion(prog, base: "MateriaACIS | None" = None) -> "MateriaACIS":
    """Vuelca una `core.programacion.Programacion` en una `MateriaACIS`, usando
    sus criterios, saberes básicos e instrumentos como punto de partida
    (pensado para la programación del curso destino, que ya está a ese nivel).
    El docente después los edita. No decide la casilla `acs_determinada`."""
    materia = base or MateriaACIS(materia=getattr(prog, "materia", "") or "")

    bloques_cri: list[str] = []
    for c in getattr(prog, "competencias", []):
        cabecera = f"Competencia específica {c.numero}".strip()
        if c.texto:
            cabecera += f" — {c.texto}"
        lineas = [cabecera]
        for cod, txt in c.criterios:
            lineas.append(f"  {cod} {txt}".rstrip())
        bloques_cri.append("\n".join(lineas))
    materia.criterios_evaluacion = "\n\n".join(bloques_cri)

    bloques_con: list[str] = []
    for bloque, items in getattr(prog, "saberes_basicos", {}).items():
        lineas = [bloque] + [f"  − {it}" for it in items]
        bloques_con.append("\n".join(lineas))
    materia.contenidos = "\n\n".join(bloques_con)

    materia.instrumentos = "\n".join(
        f"{elem}: {inst}" if elem else inst
        for elem, inst in getattr(prog, "instrumentos", [])
    )
    return materia


def datos_desde_dict(d: dict) -> DatosACIS:
    """Construye `DatosACIS` a partir de un diccionario (p. ej. leído de JSON)."""
    materias = []
    for m in d.get("materias", []):
        secu = [tuple(par) for par in m.get("secuenciacion", []) if any(par)]
        materias.append(MateriaACIS(
            materia=m.get("materia", ""),
            profesor=m.get("profesor", ""),
            departamento=m.get("departamento", ""),
            acs_determinada=bool(m.get("acs_determinada", False)),
            competencias=m.get("competencias", ""),
            criterios_evaluacion=m.get("criterios_evaluacion", ""),
            contenidos=m.get("contenidos", ""),
            metodologia=m.get("metodologia", ""),
            instrumentos=m.get("instrumentos", ""),
            unidades=m.get("unidades", ""),
            secuenciacion=secu,
        ))
    return DatosACIS(
        localidad=d.get("localidad", ""),
        fecha=d.get("fecha", ""),
        profesor_de=d.get("profesor_de", ""),
        departamento_vb=d.get("departamento_vb", ""),
        materias=materias,
    )


# --------------------------------------------------------------------------- #
# Utilidades de documento
# --------------------------------------------------------------------------- #

def _linea(valor: str, ancho: int = 40) -> str:
    """Devuelve el valor si tiene contenido; si no, una línea para rellenar."""
    return valor.strip() if valor and valor.strip() else "_" * ancho


def _sombrear(celda, fill: str) -> None:
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _texto_celda(celda, texto: str, *, negrita: bool = False, size: int = 9,
                 alineacion=None) -> None:
    celda.text = ""
    p = celda.paragraphs[0]
    if alineacion is not None:
        p.alignment = alineacion
    for i, trozo in enumerate(str(texto).split("\n")):
        par = p if i == 0 else celda.add_paragraph()
        run = par.add_run(trozo)
        run.bold = negrita
        run.font.size = Pt(size)


def _parrafo(doc, texto: str = "", *, negrita: bool = False, size: int = 10,
             alineacion=None) -> None:
    p = doc.add_paragraph()
    if alineacion is not None:
        p.alignment = alineacion
    if texto:
        run = p.add_run(texto)
        run.bold = negrita
        run.font.size = Pt(size)


def _fila_completa(tabla, texto: str, textos: dict, *, titulo: bool = False) -> None:
    """Añade una fila que ocupa las cuatro columnas."""
    fila = tabla.add_row()
    celda = fila.cells[0].merge(fila.cells[3])
    _texto_celda(celda, texto, negrita=titulo)
    if titulo:
        _sombrear(celda, textos["sombra"])


# --------------------------------------------------------------------------- #
# Bloque por materia
# --------------------------------------------------------------------------- #

def _tabla_materia(doc, materia: MateriaACIS, textos: dict) -> None:
    c = textos["cuerpo"]
    pend = textos["pendiente"]

    tabla = doc.add_table(rows=0, cols=4)
    tabla.style = "Table Grid"
    tabla.autofit = True

    _fila_completa(tabla, c["titulo"], textos, titulo=True)
    _fila_completa(tabla, f"{c['materia']} {_linea(materia.materia, 55)}", textos)

    fila = tabla.add_row()
    _texto_celda(fila.cells[0], f"{c['profesor']} {_linea(materia.profesor, 25)}")
    der = fila.cells[1].merge(fila.cells[3])
    _texto_celda(der, f"{c['departamento']} {_linea(materia.departamento, 25)}")

    _fila_completa(tabla, c["competencias_tit"], textos, titulo=True)
    _fila_completa(tabla, materia.competencias.strip() or pend, textos)

    _fila_completa(tabla, c["elementos_tit"], textos, titulo=True)

    fila = tabla.add_row()
    izq = fila.cells[0].merge(fila.cells[1])
    der = fila.cells[2].merge(fila.cells[3])
    _texto_celda(izq, c["criterios_tit"], negrita=True)
    _texto_celda(der, c["contenidos_tit"], negrita=True)

    fila = tabla.add_row()
    izq = fila.cells[0].merge(fila.cells[1])
    der = fila.cells[2].merge(fila.cells[3])
    _texto_celda(izq, materia.criterios_evaluacion.strip() or pend)
    _texto_celda(der, materia.contenidos.strip() or pend)

    _fila_completa(tabla, c["metodologia_tit"], textos, titulo=True)
    _fila_completa(tabla, materia.metodologia.strip() or pend, textos)

    _fila_completa(tabla, c["instrumentos_tit"], textos, titulo=True)
    _fila_completa(tabla, materia.instrumentos.strip() or pend, textos)

    _fila_completa(tabla, c["unidades_tit"], textos, titulo=True)
    _fila_completa(tabla, materia.unidades.strip() or pend, textos)

    _fila_completa(tabla, c["secuenciacion_tit"], textos, titulo=True)
    fila = tabla.add_row()
    ud = fila.cells[0].merge(fila.cells[2])
    _texto_celda(ud, c["ud_col"], negrita=True)
    _texto_celda(fila.cells[3], c["trimestre_col"], negrita=True)

    filas_secu = materia.secuenciacion or [("", "")]
    for unidad, trimestre in filas_secu:
        fila = tabla.add_row()
        ud = fila.cells[0].merge(fila.cells[2])
        _texto_celda(ud, _linea(unidad, 45))
        _texto_celda(fila.cells[3], _linea(trimestre, 14))


# --------------------------------------------------------------------------- #
# Bloques fijos
# --------------------------------------------------------------------------- #

def _tabla_datos_alumno(doc, textos: dict) -> None:
    d = textos["datos_alumno"]
    tabla = doc.add_table(rows=0, cols=4)
    tabla.style = "Table Grid"

    fila = tabla.add_row()
    cab = fila.cells[0].merge(fila.cells[3])
    _texto_celda(cab, d["titulo"], negrita=True, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _sombrear(cab, textos["sombra"])

    fila = tabla.add_row()
    _texto_celda(fila.cells[0], f"{d['apellidos']} {'_' * 20}")
    nombre = fila.cells[1].merge(fila.cells[2])
    _texto_celda(nombre, f"{d['nombre']} {'_' * 20}")
    _texto_celda(fila.cells[3], f"{d['nia']} {'_' * 12}")

    fila = tabla.add_row()
    _texto_celda(fila.cells[0], f"{d['fecha_nac']} {'_' * 14}")
    _texto_celda(fila.cells[1], f"{d['curso']} {'_' * 8}")
    grupo = fila.cells[2].merge(fila.cells[3])
    _texto_celda(grupo, f"{d['grupo']} {'_' * 8}")


def _bloque_firmas(doc, datos: DatosACIS, textos: dict) -> None:
    f = textos["firmas"]
    _parrafo(doc)
    _parrafo(
        doc,
        f["lugar_fecha"].format(
            localidad=_linea(datos.localidad, 25), fecha=_linea(datos.fecha, 25)
        ),
        alineacion=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    _parrafo(doc)
    p = doc.add_paragraph()
    r = p.add_run(f["profesor_de"].format(profesor_de=_linea(datos.profesor_de, 22)))
    r.font.size = Pt(10)
    r = p.add_run("\t\t\t    " + f["visto_bueno"])
    r.font.size = Pt(10)
    _parrafo(doc, f["jefe_departamento"], alineacion=WD_ALIGN_PARAGRAPH.RIGHT)
    _parrafo(
        doc,
        f["jefe_departamento_de"].format(departamento_vb=_linea(datos.departamento_vb, 22)),
        alineacion=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    for _ in range(3):
        _parrafo(doc)
    p = doc.add_paragraph()
    r = p.add_run(f["fdo"] + "_" * 24 + "\t\t\t\t\t" + f["fdo"] + "_" * 24)
    r.font.size = Pt(10)
    for _ in range(4):
        _parrafo(doc)
    _parrafo(doc, f["jefe_estudios"], negrita=True)
    _parrafo(doc, f["expediente"], size=8)


def _poner_pie(doc, texto: str) -> None:
    parrafo = doc.sections[0].footer.paragraphs[0]
    parrafo.text = texto
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in parrafo.runs:
        run.font.size = Pt(8)


# --------------------------------------------------------------------------- #
# Punto de entrada
# --------------------------------------------------------------------------- #

def generar_acis(ruta_salida: str, datos: DatosACIS) -> tuple[str, list[str]]:
    """Crea en `ruta_salida` el Anexo III.b con los datos aportados.

    Devuelve `(ruta, avisos)`. `avisos` recoge las materias que se han dejado
    fuera por no constar la determinación de ACS por el equipo de orientación.
    """
    textos = cargar_textos()
    doc = Document()

    seccion = doc.sections[0]
    seccion.page_width = Cm(21)
    seccion.page_height = Cm(29.7)
    seccion.left_margin = Cm(3)
    seccion.right_margin = Cm(3)
    seccion.top_margin = Cm(3.5)
    seccion.bottom_margin = Cm(1.75)

    _poner_pie(doc, textos["footer"])

    _parrafo(doc, textos["anexo"], negrita=True, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _parrafo(doc, textos["titulo"], negrita=True, alineacion=WD_ALIGN_PARAGRAPH.CENTER)

    aviso = doc.add_paragraph()
    run = aviso.add_run(textos["marca_borrador"])
    run.italic = True
    run.bold = True
    run.font.size = Pt(9)

    _parrafo(doc)
    _tabla_datos_alumno(doc, textos)

    avisos: list[str] = []
    for materia in datos.materias:
        if not materia.acs_determinada:
            avisos.append(textos["aviso_sin_acs"].format(
                materia=materia.materia.strip() or "materia sin nombre"))
            continue
        _parrafo(doc)
        _tabla_materia(doc, materia, textos)

    if avisos:
        _parrafo(doc)
        p = doc.add_paragraph()
        r = p.add_run("Materias no incluidas:")
        r.bold = True
        r.font.size = Pt(9)
        for texto in avisos:
            item = doc.add_paragraph(texto, style="List Bullet")
            for run in item.runs:
                run.font.size = Pt(9)

    _bloque_firmas(doc, datos, textos)

    carpeta = os.path.dirname(os.path.abspath(ruta_salida))
    os.makedirs(carpeta, exist_ok=True)
    doc.save(ruta_salida)
    return ruta_salida, avisos
