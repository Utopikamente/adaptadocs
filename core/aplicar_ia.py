"""Vuelca el resultado de la IA (`core.ia.adaptar_contenido`) sobre un
`Document` abierto: reescribe párrafos, crea listas de pasos y añade las
secciones de resumen, glosario y preguntas.
"""

from __future__ import annotations

from typing import Callable

from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


# --------------------------------------------------------------------------- #
# Utilidades de edición
# --------------------------------------------------------------------------- #

def _reemplazar_texto(parrafo: Paragraph, texto: str) -> None:
    """Cambia el texto del párrafo conservando el formato del primer run."""
    if parrafo.runs:
        parrafo.runs[0].text = texto
        for run in parrafo.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        parrafo.add_run(texto)


def _insertar_despues(parrafo: Paragraph, texto: str = "", estilo: str | None = None) -> Paragraph:
    nuevo_el = OxmlElement("w:p")
    parrafo._p.addnext(nuevo_el)
    nuevo = Paragraph(nuevo_el, parrafo._parent)
    if estilo:
        try:
            nuevo.style = estilo
        except KeyError:
            pass
    if texto:
        nuevo.add_run(texto)
    return nuevo


def _estilo_disponible(doc, *nombres: str) -> str | None:
    for nombre in nombres:
        try:
            doc.styles[nombre]
            return nombre
        except KeyError:
            continue
    return None


# --------------------------------------------------------------------------- #
# Secciones nuevas
# --------------------------------------------------------------------------- #

def _anadir_glosario(doc, glosario: list[dict]) -> int:
    doc.add_heading("Glosario", level=1)
    for entrada in glosario:
        termino = str(entrada.get("termino", "")).strip()
        definicion = str(entrada.get("definicion", "")).strip()
        if not termino:
            continue
        p = doc.add_paragraph()
        p.add_run(f"{termino}: ").bold = True
        p.add_run(definicion)
    return len(glosario)


def _anadir_preguntas(doc, preguntas: list[str]) -> int:
    doc.add_heading("Preguntas de comprensión", level=1)
    estilo = _estilo_disponible(doc, "List Number") or "Normal"
    for i, pregunta in enumerate(preguntas, start=1):
        texto = str(pregunta).strip()
        if not texto:
            continue
        if estilo == "Normal":
            doc.add_paragraph(f"{i}. {texto}")
        else:
            doc.add_paragraph(texto, style=estilo)
    return len(preguntas)


def _insertar_resumen(doc, bloques_resumen: list[dict]) -> int:
    if not doc.paragraphs:
        return 0
    ancla = _insertar_despues(
        doc.paragraphs[0],
        "Resumen",
        estilo=_estilo_disponible(doc, "Heading 2", "Título 2", "Heading 1"),
    )
    estilo_punto = _estilo_disponible(doc, "List Bullet")
    total = 0
    for bloque in bloques_resumen:
        apartado = str(bloque.get("apartado", "")).strip()
        if apartado:
            ancla = _insertar_despues(ancla, apartado)
            if ancla.runs:
                ancla.runs[0].bold = True
        for punto in bloque.get("puntos", []):
            texto = str(punto).strip()
            if not texto:
                continue
            ancla = _insertar_despues(ancla, texto if estilo_punto else f"• {texto}", estilo_punto)
            total += 1
    return total


# --------------------------------------------------------------------------- #
# Punto de entrada
# --------------------------------------------------------------------------- #

def aplicar_resultado(
    doc,
    datos: dict,
    parrafos_por_id: dict[int, Paragraph],
    registrar: Callable[[str], None] = lambda mensaje: None,
) -> dict:
    """Aplica `datos` (salida de `core.ia`) sobre `doc`. Devuelve un resumen."""

    resumen = {
        "simplificados": 0,
        "en_pasos": 0,
        "glosario": 0,
        "preguntas": 0,
        "resumen": 0,
        "preguntas_divididas": 0,
    }

    ids_en_pasos = {int(x["id"]) for x in datos.get("parrafos_en_pasos", []) if "id" in x}
    ids_preguntas_divididas = {
        int(x["id"]) for x in datos.get("preguntas_divididas", []) if "id" in x
    }

    # 1) Reescritura de párrafos
    for item in datos.get("parrafos_simplificados", []):
        try:
            pid = int(item["id"])
        except (KeyError, ValueError, TypeError):
            continue
        if pid in ids_en_pasos or pid in ids_preguntas_divididas:
            continue
        parrafo = parrafos_por_id.get(pid)
        texto = str(item.get("texto", "")).strip()
        if parrafo is not None and texto:
            _reemplazar_texto(parrafo, texto)
            resumen["simplificados"] += 1

    # 2) Párrafos convertidos en listas de pasos
    estilo_pasos = _estilo_disponible(doc, "List Number")
    for item in datos.get("parrafos_en_pasos", []):
        try:
            pid = int(item["id"])
        except (KeyError, ValueError, TypeError):
            continue
        if pid in ids_preguntas_divididas:
            continue
        parrafo = parrafos_por_id.get(pid)
        pasos = [str(s).strip() for s in item.get("pasos", []) if str(s).strip()]
        if parrafo is None or not pasos:
            continue
        _reemplazar_texto(parrafo, pasos[0] if not estilo_pasos else pasos[0])
        if estilo_pasos:
            try:
                parrafo.style = estilo_pasos
            except KeyError:
                pass
        ancla = parrafo
        for paso in pasos[1:]:
            ancla = _insertar_despues(ancla, paso, estilo_pasos)
        resumen["en_pasos"] += 1

    # 2b) Preguntas compuestas divididas en varias más cortas
    for item in datos.get("preguntas_divididas", []):
        try:
            pid = int(item["id"])
        except (KeyError, ValueError, TypeError):
            continue
        parrafo = parrafos_por_id.get(pid)
        subpreguntas = [str(s).strip() for s in item.get("subpreguntas", []) if str(s).strip()]
        if parrafo is None or not subpreguntas:
            continue
        _reemplazar_texto(parrafo, subpreguntas[0])
        ancla = parrafo
        for subpregunta in subpreguntas[1:]:
            ancla = _insertar_despues(ancla, subpregunta)
        resumen["preguntas_divididas"] += 1

    # 3) Resumen al principio (antes de tocar el final del documento)
    if datos.get("resumen"):
        resumen["resumen"] = _insertar_resumen(doc, datos["resumen"])
        registrar("Resumen añadido al principio.")

    # 4) Glosario y preguntas al final
    if datos.get("glosario"):
        resumen["glosario"] = _anadir_glosario(doc, datos["glosario"])
        registrar("Glosario añadido al final.")
    if datos.get("preguntas"):
        resumen["preguntas"] = _anadir_preguntas(doc, datos["preguntas"])
        registrar("Preguntas de comprensión añadidas al final.")

    return resumen
