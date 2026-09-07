"""Genera un documento interno con las adaptaciones aplicadas a una prueba.

Es una hoja de consulta rápida para el profesorado: si se pregunta qué se ha
adaptado en un material, este documento lo resume. **No** sustituye a los
anexos oficiales que van al expediente académico (Anexo VI en ESO, Anexo III
de la Orden 130/2023 en Primaria).

Los textos fijos y las plantillas de cada línea están en `registro.json`, que
se puede editar sin tocar el código.
"""

from __future__ import annotations

import json
import os
from datetime import date
from typing import Any

from docx import Document
from docx.shared import Pt

from .transformador import OpcionesAdaptacion

_RUTA_DATOS = os.path.join(os.path.dirname(__file__), "registro.json")


def cargar_textos() -> dict:
    """Lee `registro.json` (textos y plantillas editables)."""
    with open(_RUTA_DATOS, encoding="utf-8") as f:
        return json.load(f)


def ruta_registro(ruta_documento_adaptado: str) -> str:
    """A partir de «X (adaptado).docx» devuelve «X (adaptado) — registro.docx»."""
    carpeta, nombre = os.path.split(ruta_documento_adaptado)
    raiz, _ = os.path.splitext(nombre)
    return os.path.join(carpeta, f"{raiz} — registro.docx")


def _num(valor: Any) -> str:
    """Número legible en español: 14.0 -> «14»; 1.5 -> «1,5»."""
    f = float(valor)
    if f == int(f):
        return str(int(f))
    return str(f).replace(".", ",")


def _formatea(plantilla: str | None, **valores: Any) -> str | None:
    if not plantilla:
        return None
    try:
        return plantilla.format(**valores)
    except (KeyError, IndexError):
        return plantilla


def lineas_formato(opciones: OpcionesAdaptacion, resumen: dict, textos: dict) -> list[str]:
    """Lista de frases que describen las adaptaciones de formato aplicadas."""
    p = textos["lineas_formato"]
    fuente = (opciones.fuente or "").strip()
    tamano = opciones.tamano_pt
    out: list[str | None] = []

    if fuente and tamano:
        out.append(_formatea(p.get("tipografia"), fuente=fuente, tamano=_num(tamano)))
    elif fuente:
        out.append(_formatea(p.get("solo_fuente"), fuente=fuente))
    elif tamano:
        out.append(_formatea(p.get("solo_tamano"), tamano=_num(tamano)))

    if opciones.interlineado:
        out.append(_formatea(p.get("interlineado"), interlineado=_num(opciones.interlineado)))
    if opciones.espacio_despues_pt:
        out.append(_formatea(p.get("espacio_parrafo"), espacio=_num(opciones.espacio_despues_pt)))
    if opciones.alinear_izquierda:
        out.append(_formatea(p.get("alinear_izquierda")))
    if opciones.una_columna:
        out.append(_formatea(p.get("una_columna")))
    if opciones.margenes_cm is not None:
        out.append(_formatea(p.get("margenes"), margenes=_num(opciones.margenes_cm)))
    if opciones.alto_contraste:
        out.append(_formatea(p.get("alto_contraste")))
    if opciones.negrita_titulos:
        out.append(_formatea(p.get("negrita_titulos")))
    if opciones.convertir_vinetas_en_pasos and resumen.get("vinetas_convertidas"):
        out.append(_formatea(p.get("vinetas_numeradas")))
    if resumen.get("procedimientos_en_pasos"):
        out.append(_formatea(p.get("pasos"), procedimientos=resumen["procedimientos_en_pasos"]))
    if resumen.get("resaltados"):
        out.append(_formatea(p.get("resaltado"), resaltados=resumen["resaltados"]))
    if resumen.get("preguntas_numeradas"):
        out.append(_formatea(p.get("preguntas_numeradas"),
                             preguntas_numeradas=resumen["preguntas_numeradas"]))
    if resumen.get("preguntas_con_espacio"):
        out.append(_formatea(p.get("espacio_respuestas"),
                             preguntas=resumen["preguntas_con_espacio"],
                             lineas=opciones.espacio_respuestas))

    return [x for x in out if x]


def lineas_ia(resumen_ia: dict, opciones_ia: Any, textos: dict) -> list[str]:
    """Lista de frases que describen lo que hizo la capa de IA, si se usó."""
    if not resumen_ia or opciones_ia is None:
        return []
    p = textos["lineas_ia"]
    out: list[str | None] = []
    if getattr(opciones_ia, "simplificar", False):
        out.append(_formatea(p.get("simplificar"), nivel=getattr(opciones_ia, "nivel", "")))
    if getattr(opciones_ia, "pasos", False):
        out.append(_formatea(p.get("pasos")))
    if getattr(opciones_ia, "glosario", False):
        out.append(_formatea(p.get("glosario")))
    if getattr(opciones_ia, "resumen", False):
        out.append(_formatea(p.get("resumen")))
    if getattr(opciones_ia, "preguntas", False):
        out.append(_formatea(p.get("preguntas"), preguntas=resumen_ia.get("preguntas", 0)))
    if getattr(opciones_ia, "dividir_preguntas", False):
        out.append(_formatea(p.get("dividir_preguntas"),
                             divididas=resumen_ia.get("preguntas_divididas", 0)))
    return [x for x in out if x]


def generar_registro(
    ruta_salida: str,
    nombre_documento: str,
    opciones: OpcionesAdaptacion,
    resumen: dict,
    resumen_ia: dict | None = None,
    opciones_ia: Any = None,
    fecha: date | None = None,
) -> str:
    """Crea en `ruta_salida` un .docx con las adaptaciones aplicadas a
    `nombre_documento`. Devuelve la ruta del archivo escrito."""
    textos = cargar_textos()
    doc = Document()

    cabecera = doc.add_paragraph()
    run = cabecera.add_run(textos["titulo"])
    run.bold = True
    run.font.size = Pt(14)

    sub = doc.add_paragraph(textos["subtitulo"])
    for r in sub.runs:
        r.italic = True
        r.font.size = Pt(9)

    doc.add_paragraph(f"{textos['etiqueta_documento']} {nombre_documento}")
    doc.add_paragraph(f"{textos['etiqueta_fecha']} {(fecha or date.today()).strftime('%d/%m/%Y')}")
    doc.add_paragraph(f"{textos['etiqueta_alumno']} " + "_" * 45)
    doc.add_paragraph(f"{textos['etiqueta_barreras']} " + "_" * 30)

    doc.add_paragraph()
    doc.add_paragraph(textos["marco_legal"])

    titulo_fmt = doc.add_paragraph()
    titulo_fmt.add_run(textos["titulo_formato"]).bold = True
    formato = lineas_formato(opciones, resumen, textos)
    if formato:
        for linea in formato:
            doc.add_paragraph(linea, style="List Bullet")
    else:
        doc.add_paragraph(textos["sin_adaptaciones"])

    ia = lineas_ia(resumen_ia or {}, opciones_ia, textos)
    if ia:
        titulo_ia = doc.add_paragraph()
        titulo_ia.add_run(textos["titulo_ia"]).bold = True
        for linea in ia:
            doc.add_paragraph(linea, style="List Bullet")

    doc.add_paragraph()
    nota = doc.add_paragraph(textos["nota_pie"])
    for r in nota.runs:
        r.italic = True
        r.font.size = Pt(8)

    carpeta = os.path.dirname(os.path.abspath(ruta_salida))
    os.makedirs(carpeta, exist_ok=True)
    doc.save(ruta_salida)
    return ruta_salida
