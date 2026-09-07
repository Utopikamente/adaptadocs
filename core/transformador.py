"""Transformaciones de formato sobre documentos .docx.

Todo lo que hay aquí es determinista y sin conexión.
"""

from __future__ import annotations

import os
import re
from copy import deepcopy
from dataclasses import dataclass, field
from typing import Callable, Iterable, Iterator

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run


# --------------------------------------------------------------------------- #
# Opciones
# --------------------------------------------------------------------------- #

COLORES_RESALTADO = {
    "AMARILLO": WD_COLOR_INDEX.YELLOW,
    "VERDE": WD_COLOR_INDEX.BRIGHT_GREEN,
    "TURQUESA": WD_COLOR_INDEX.TURQUOISE,
    "ROSA": WD_COLOR_INDEX.PINK,
    "GRIS": WD_COLOR_INDEX.GRAY_25,
}


@dataclass
class OpcionesAdaptacion:
    """Conjunto de ajustes de formato que se aplican a un documento."""

    # Tipografía
    fuente: str | None = "Verdana"
    tamano_pt: float | None = 14.0

    # Espaciado
    interlineado: float | None = 1.5
    espacio_despues_pt: float | None = 24.0

    # Disposición
    alinear_izquierda: bool = True          # convierte el texto justificado en alineado a la izquierda
    margenes_cm: float | None = 2.5         # None = no tocar los márgenes
    una_columna: bool = True                # fuerza una sola columna
    alto_contraste: bool = True             # pone todo el texto en negro

    # Ayudas visuales
    negrita_titulos: bool = True
    convertir_vinetas_en_pasos: bool = False
    # Separar en pasos numerados los párrafos que describen un procedimiento.
    #   "no"        -> no hacer nada
    #   "marcados"  -> solo los que empiecen por la marca (p. ej. «PASOS:»)
    #   "auto"      -> detectar por heurística (conectores de secuencia, verbos)
    separar_en_pasos: str = "no"
    marca_pasos: str = "PASOS:"
    resaltar_palabras: list[str] = field(default_factory=list)
    color_resaltado: str = "AMARILLO"
    # Preguntas de examen/ficha: se detectan los párrafos que terminan en «?».
    numerar_preguntas: bool = False   # renumera esos párrafos como lista numerada
    espacio_respuestas: int = 0       # líneas en blanco tras cada pregunta (0 = ninguna)

    def copia(self) -> "OpcionesAdaptacion":
        return deepcopy(self)


# --------------------------------------------------------------------------- #
# Utilidades de recorrido
# --------------------------------------------------------------------------- #

def _iter_parrafos(contenedor) -> Iterator[Paragraph]:
    """Devuelve todos los párrafos de un contenedor (documento, celda,
    cabecera o pie), entrando también en las tablas anidadas."""
    for parrafo in contenedor.paragraphs:
        yield parrafo
    for tabla in contenedor.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                yield from _iter_parrafos(celda)


def _contenedores(doc) -> Iterable:
    yield doc
    for seccion in doc.sections:
        yield seccion.header
        yield seccion.first_page_header
        yield seccion.even_page_header
        yield seccion.footer
        yield seccion.first_page_footer
        yield seccion.even_page_footer


# --------------------------------------------------------------------------- #
# Fuente
# --------------------------------------------------------------------------- #

def _forzar_rfonts(rpr, fuente: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), fuente)


def _formatear_run(run: Run, o: OpcionesAdaptacion) -> None:
    if o.fuente:
        run.font.name = o.fuente
        _forzar_rfonts(run._element.get_or_add_rPr(), o.fuente)
    if o.tamano_pt:
        run.font.size = Pt(o.tamano_pt)
    if o.alto_contraste:
        run.font.color.rgb = RGBColor(0, 0, 0)


# --------------------------------------------------------------------------- #
# Párrafo
# --------------------------------------------------------------------------- #

_PREFIJOS_TITULO = ("heading", "título", "titulo", "subtitle", "subtítulo")
_MARCAS_VINETA = ("bullet", "viñeta", "vineta", "list bullet")


def _es_titulo(parrafo: Paragraph) -> bool:
    nombre = (parrafo.style.name or "").lower() if parrafo.style else ""
    return any(nombre.startswith(p) for p in _PREFIJOS_TITULO)


def _es_vineta(parrafo: Paragraph) -> bool:
    nombre = (parrafo.style.name or "").lower() if parrafo.style else ""
    return any(m in nombre for m in _MARCAS_VINETA)


def _formatear_parrafo(parrafo: Paragraph, o: OpcionesAdaptacion) -> None:
    pf = parrafo.paragraph_format
    if o.interlineado:
        pf.line_spacing = float(o.interlineado)
    if o.espacio_despues_pt is not None:
        pf.space_after = Pt(o.espacio_despues_pt)
    if o.alinear_izquierda and parrafo.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT


# --------------------------------------------------------------------------- #
# Procedimientos -> pasos numerados (sin IA, determinista)
# --------------------------------------------------------------------------- #

_CONECTORES_SECUENCIA = (
    r"luego", r"despu[eé]s", r"a continuaci[oó]n", r"entonces", r"seguidamente",
    r"acto seguido", r"m[aá]s tarde", r"por [uú]ltimo", r"finalmente", r"tras esto",
    r"en primer lugar", r"en segundo lugar", r"en tercer lugar", r"primero",
)
_RE_CONECTOR = re.compile(
    r"\s*[,;.]?\s*\b(?:" + "|".join(_CONECTORES_SECUENCIA) + r")\b[\s,]*",
    re.IGNORECASE,
)
_RE_NUMERADO = re.compile(r"(?:^|\s)\d{1,2}[.)]\s+")
_RE_FIN_FRASE = re.compile(r"(?<=[.;:])\s+")

_MARCAS_SECUENCIA = (
    "primero", "en primer lugar", "en segundo lugar", "en tercer lugar",
    "luego", "después", "despues", "a continuación", "a continuacion",
    "seguidamente", "acto seguido", "por último", "por ultimo", "finalmente",
    "paso 1", "paso 2", "1.", "2.", "3.",
)
_VERBOS_INSTRUCCION = {
    "abre", "pulsa", "haz", "ve", "coge", "escribe", "marca", "selecciona",
    "elige", "escoge", "guarda", "cierra", "copia", "pega", "dibuja", "colorea",
    "lee", "rodea", "subraya", "tacha", "completa", "une", "ordena", "clasifica",
    "calcula", "resuelve", "observa", "anota", "repasa", "recorta", "relaciona",
    "identifica", "señala", "senala", "indica", "busca", "corrige", "comprueba",
    "repite", "suma", "resta", "multiplica", "divide", "traza", "rellena",
    "contesta", "responde", "explica", "describe", "nombra", "enumera",
}


def _re_marca_pasos(marca: str) -> re.Pattern:
    base = re.escape(marca.strip().rstrip(":.-–").strip())
    return re.compile(r"^\s*" + base + r"\s*[:.\-–]\s*", re.IGNORECASE)


def _parece_procedimiento(texto: str) -> bool:
    bajo = texto.lower()
    if sum(1 for m in _MARCAS_SECUENCIA if m in bajo) >= 2:
        return True
    frases = [f.strip() for f in _RE_FIN_FRASE.split(texto) if f.strip()]
    if len(frases) >= 3:
        imperativas = 0
        for frase in frases:
            primera = re.split(r"[^0-9A-Za-zÁÉÍÓÚÑáéíóúñ]+", frase.strip(), maxsplit=1)[0]
            if primera.lower() in _VERBOS_INSTRUCCION:
                imperativas += 1
        if imperativas >= max(2, len(frases) // 2):
            return True
    return False


def _partir_en_pasos(texto: str) -> list[str]:
    """Trocea un texto en pasos. No reescribe: solo corta y limpia."""
    t = re.sub(r"\s+", " ", texto).strip()

    numerado = [p.strip() for p in _RE_NUMERADO.split(t) if p.strip()]
    if len(numerado) >= 2:
        crudos = numerado
    else:
        t = _RE_CONECTOR.sub("\n", t)
        t = _RE_FIN_FRASE.sub("\n", t)
        crudos = [p for p in t.split("\n")]

    pasos: list[str] = []
    for trozo in crudos:
        p = trozo.strip(" .;:,\t")
        if len(p) < 3:
            continue
        p = p[0].upper() + p[1:]
        if p[-1] not in ".!?":
            p += "."
        pasos.append(p)
    return pasos


def _abstract_id_lista_numerada(doc):
    """Devuelve el `abstractNumId` que usa el estilo «List Number», o None."""
    try:
        estilo = doc.styles["List Number"].element
    except KeyError:
        return None
    vals = estilo.xpath(".//w:numPr/w:numId/@w:val")
    if not vals:
        return None
    try:
        numbering = doc.part.numbering_part.element
        num = numbering.num_having_numId(int(vals[0]))
    except (KeyError, ValueError, NotImplementedError):
        return None
    return num.abstractNumId.val


def _nueva_lista_reiniciada(doc, abstract_id):
    """Crea una numeración nueva con el mismo formato que «List Number» pero
    que vuelve a empezar en 1. Devuelve su `numId`, o None si no se puede."""
    if abstract_id is None:
        return None
    try:
        numbering = doc.part.numbering_part.element
        num = numbering.add_num(abstract_id)
        num.add_lvlOverride(0).add_startOverride(1)
        return num.numId
    except (KeyError, ValueError, NotImplementedError):
        return None


def _asignar_numeracion(parrafo: Paragraph, num_id: int) -> None:
    """Fuerza el `numId` (la lista concreta) de un párrafo, sin depender de
    la numeración global del estilo."""
    numPr = parrafo._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = num_id


def _separar_procedimientos(doc, o: OpcionesAdaptacion) -> int:
    """Convierte en listas numeradas los párrafos de primer nivel que sean
    (o estén marcados como) un procedimiento. Cada procedimiento estrena su
    propia numeración, que empieza en 1. Devuelve cuántos ha convertido."""
    if o.separar_en_pasos not in ("marcados", "auto"):
        return 0

    re_marca = _re_marca_pasos(o.marca_pasos)
    abstract_id = _abstract_id_lista_numerada(doc)
    convertidos = 0

    for parrafo in list(doc.paragraphs):
        if _es_titulo(parrafo) or _es_vineta(parrafo):
            continue
        texto = parrafo.text.strip()
        if not texto:
            continue

        m = re_marca.match(texto)
        if m:
            cuerpo = texto[m.end():].strip()
        elif o.separar_en_pasos == "auto" and _parece_procedimiento(texto):
            cuerpo = texto
        else:
            continue

        pasos = _partir_en_pasos(cuerpo)
        if len(pasos) < 2:
            continue

        _reemplazar_texto_parrafo(parrafo, pasos[0])
        try:
            parrafo.style = doc.styles["List Number"]
        except KeyError:
            pass
        num_id = _nueva_lista_reiniciada(doc, abstract_id)
        if num_id is not None:
            _asignar_numeracion(parrafo, num_id)
        ancla = parrafo
        for paso in pasos[1:]:
            ancla = _nuevo_parrafo_despues(ancla, paso, "List Number")
            if num_id is not None:
                _asignar_numeracion(ancla, num_id)
        convertidos += 1

    return convertidos


def _reemplazar_texto_parrafo(parrafo: Paragraph, texto: str) -> None:
    if parrafo.runs:
        parrafo.runs[0].text = texto
        for run in parrafo.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        parrafo.add_run(texto)


def _nuevo_parrafo_despues(parrafo: Paragraph, texto: str, estilo: str | None) -> Paragraph:
    nuevo_el = OxmlElement("w:p")
    parrafo._p.addnext(nuevo_el)
    nuevo = Paragraph(nuevo_el, parrafo._parent)
    if estilo:
        try:
            nuevo.style = estilo
        except KeyError:
            pass
    nuevo.add_run(texto)
    return nuevo


# --------------------------------------------------------------------------- #
# Preguntas: numerarlas y dejar espacio para responder
# --------------------------------------------------------------------------- #

def _es_pregunta(texto: str) -> bool:
    """Un párrafo se considera pregunta si termina en «?» (admite un cierre
    de comilla o paréntesis después del interrogante)."""
    t = texto.strip().rstrip("»\"')]")
    return t.endswith("?")


def _procesar_preguntas(doc, o: OpcionesAdaptacion) -> dict:
    """Detecta los párrafos que son preguntas y, según las opciones, los
    numera de forma consecutiva y/o añade líneas en blanco para responder."""
    resumen = {"preguntas_numeradas": 0, "preguntas_con_espacio": 0}
    if not o.numerar_preguntas and not o.espacio_respuestas:
        return resumen

    for parrafo in list(doc.paragraphs):
        if _es_titulo(parrafo):
            continue
        texto = parrafo.text.strip()
        if not texto or not _es_pregunta(texto):
            continue

        if o.espacio_respuestas > 0:
            ancla = parrafo
            for _ in range(o.espacio_respuestas):
                ancla = _nuevo_parrafo_despues(ancla, "", None)
            resumen["preguntas_con_espacio"] += 1

        if o.numerar_preguntas:
            try:
                parrafo.style = doc.styles["List Number"]
                resumen["preguntas_numeradas"] += 1
            except KeyError:
                pass

    return resumen


# --------------------------------------------------------------------------- #
# Resaltado de palabras clave
# --------------------------------------------------------------------------- #

def _clonar_run(base: Run, texto: str, ancla: Run) -> Run:
    nuevo_el = deepcopy(base._element)
    ancla._element.addnext(nuevo_el)
    nuevo = Run(nuevo_el, base._parent)
    nuevo.text = texto
    return nuevo


def _resaltar_en_parrafo(parrafo: Paragraph, patron: re.Pattern, color) -> int:
    """Divide los runs para poder marcar solo las palabras buscadas.
    Devuelve el número de coincidencias resaltadas."""
    total = 0
    for run in list(parrafo.runs):
        texto = run.text
        if not texto:
            continue
        trozos = patron.split(texto)
        if len(trozos) == 1:
            continue
        run.text = trozos[0]
        ancla = run
        for i, trozo in enumerate(trozos[1:], start=1):
            if trozo == "":
                continue
            nuevo = _clonar_run(run, trozo, ancla)
            if i % 2 == 1:                       # los índices impares son coincidencias
                nuevo.font.highlight_color = color
                total += 1
            else:
                nuevo.font.highlight_color = None
            ancla = nuevo
    return total


def _compilar_patron(palabras: list[str]) -> re.Pattern | None:
    limpias = [p.strip() for p in palabras if p and p.strip()]
    if not limpias:
        return None
    alternativas = "|".join(re.escape(p) for p in sorted(limpias, key=len, reverse=True))
    return re.compile(f"({alternativas})", re.IGNORECASE)


# --------------------------------------------------------------------------- #
# Secciones (márgenes y columnas)
# --------------------------------------------------------------------------- #

def _una_columna(seccion) -> None:
    sectPr = seccion._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        return
    cols.set(qn("w:num"), "1")
    for hijo in list(cols):
        cols.remove(hijo)


def _ajustar_seccion(seccion, o: OpcionesAdaptacion) -> None:
    if o.margenes_cm is not None:
        margen = Cm(o.margenes_cm)
        seccion.left_margin = margen
        seccion.right_margin = margen
        seccion.top_margin = margen
        seccion.bottom_margin = margen
    if o.una_columna:
        _una_columna(seccion)


# --------------------------------------------------------------------------- #
# Estilo Normal
# --------------------------------------------------------------------------- #

def _ajustar_estilo_normal(doc, o: OpcionesAdaptacion) -> None:
    try:
        normal = doc.styles["Normal"]
    except KeyError:
        return
    if o.fuente:
        normal.font.name = o.fuente
        _forzar_rfonts(normal.element.get_or_add_rPr(), o.fuente)
    if o.tamano_pt:
        normal.font.size = Pt(o.tamano_pt)
    pf = normal.paragraph_format
    if o.interlineado:
        pf.line_spacing = float(o.interlineado)
    if o.espacio_despues_pt is not None:
        pf.space_after = Pt(o.espacio_despues_pt)


# --------------------------------------------------------------------------- #
# Punto de entrada
# --------------------------------------------------------------------------- #

def _validar_rutas(ruta_entrada: str, ruta_salida: str) -> None:
    ext = os.path.splitext(ruta_entrada)[1].lower()
    if ext == ".doc":
        raise ValueError(
            "El formato .doc (Word 97-2003) no es compatible. "
            "Abre el archivo en Word y guárdalo como .docx."
        )
    if ext != ".docx":
        raise ValueError(f"Se esperaba un archivo .docx y se recibió «{ext or 'sin extensión'}».")
    if os.path.abspath(ruta_entrada) == os.path.abspath(ruta_salida):
        raise ValueError("El archivo de salida no puede ser el mismo que el de entrada.")


def aplicar_formato(
    doc,
    opciones: OpcionesAdaptacion,
    registrar: Callable[[str], None] = lambda mensaje: None,
) -> dict:
    """Aplica las opciones de formato sobre un `Document` ya abierto.
    Devuelve un resumen de lo hecho. No guarda el archivo."""

    resumen = {
        "parrafos": 0, "runs": 0, "resaltados": 0,
        "vinetas_convertidas": 0, "procedimientos_en_pasos": 0,
        "preguntas_numeradas": 0, "preguntas_con_espacio": 0,
    }

    _ajustar_estilo_normal(doc, opciones)
    for seccion in doc.sections:
        _ajustar_seccion(seccion, opciones)

    # Antes de formatear: trocear procedimientos en pasos y tratar las
    # preguntas (crean párrafos nuevos que el bucle de abajo recogerá al
    # releer doc.paragraphs).
    resumen["procedimientos_en_pasos"] = _separar_procedimientos(doc, opciones)
    resumen.update(_procesar_preguntas(doc, opciones))

    patron = _compilar_patron(opciones.resaltar_palabras)
    color = COLORES_RESALTADO.get(opciones.color_resaltado.upper(), WD_COLOR_INDEX.YELLOW)

    for contenedor in _contenedores(doc):
        for parrafo in _iter_parrafos(contenedor):
            resumen["parrafos"] += 1

            if opciones.convertir_vinetas_en_pasos and _es_vineta(parrafo):
                try:
                    parrafo.style = doc.styles["List Number"]
                    resumen["vinetas_convertidas"] += 1
                except KeyError:
                    pass

            _formatear_parrafo(parrafo, opciones)

            es_titulo = opciones.negrita_titulos and _es_titulo(parrafo)
            for run in parrafo.runs:
                _formatear_run(run, opciones)
                if es_titulo:
                    run.bold = True
                resumen["runs"] += 1

            if patron is not None:
                resumen["resaltados"] += _resaltar_en_parrafo(parrafo, patron, color)

    return resumen


def adaptar_documento(
    ruta_entrada: str,
    ruta_salida: str,
    opciones: OpcionesAdaptacion,
    registrar: Callable[[str], None] = lambda mensaje: None,
) -> dict:
    """Lee `ruta_entrada`, aplica `opciones` de formato y guarda el resultado
    en `ruta_salida`. Devuelve un pequeño resumen de lo que se ha hecho."""

    _validar_rutas(ruta_entrada, ruta_salida)

    registrar(f"Abriendo «{os.path.basename(ruta_entrada)}»…")
    doc = Document(ruta_entrada)

    resumen = aplicar_formato(doc, opciones, registrar)

    carpeta = os.path.dirname(os.path.abspath(ruta_salida))
    os.makedirs(carpeta, exist_ok=True)
    doc.save(ruta_salida)
    registrar(f"Guardado en «{ruta_salida}».")

    return resumen
