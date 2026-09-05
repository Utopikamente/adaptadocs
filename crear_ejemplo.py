"""Genera un documento de ejemplo para probar Adaptadocs.

    python crear_ejemplo.py
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def main() -> None:
    doc = Document()

    doc.add_heading("La fotosíntesis", level=1)

    p = doc.add_paragraph(
        "La fotosíntesis es el proceso por el cual las plantas fabrican su propio "
        "alimento. Utilizan la luz del sol, el agua del suelo y el dióxido de carbono "
        "del aire. Como resultado, producen glucosa y liberan oxígeno. Este proceso "
        "es muy importante para la vida en la Tierra."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_heading("Pasos del proceso", level=2)
    for paso in [
        "La planta absorbe agua por las raíces.",
        "Las hojas captan la luz del sol.",
        "El dióxido de carbono entra por los estomas.",
        "Se produce glucosa y se libera oxígeno.",
    ]:
        doc.add_paragraph(paso, style="List Bullet")

    doc.add_heading("Vocabulario importante", level=2)
    p2 = doc.add_paragraph(
        "Recuerda estas palabras: glucosa, oxígeno, dióxido de carbono y estomas. "
        "La glucosa es el alimento de la planta."
    )
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_heading("Actividad", level=2)
    doc.add_paragraph(
        "PASOS: Coge una planta pequeña. Ponla cerca de la ventana. "
        "Riégala con un poco de agua. Observa las hojas cada día durante una semana."
    )
    doc.add_paragraph(
        "Primero, dibuja una planta en tu cuaderno. Luego colorea las hojas de verde. "
        "Después escribe debajo la palabra fotosíntesis. Por último, enseña el dibujo "
        "a un compañero."
    )

    doc.add_heading("Preguntas", level=2)
    doc.add_paragraph("¿Qué necesita una planta para hacer la fotosíntesis?")
    doc.add_paragraph("¿Qué produce la planta como resultado de este proceso?")

    doc.save("ejemplo.docx")
    print("Creado ejemplo.docx")


if __name__ == "__main__":
    main()
