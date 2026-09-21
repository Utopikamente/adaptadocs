"""Comprobación rápida de que la adaptación hace lo esperado.

    python prueba.py

Genera un documento de ejemplo, lo adapta con varios perfiles y verifica el
resultado. Sale con código 0 si todo va bien, 1 si algo falla. Lo usa también
la integración continua.
"""

from __future__ import annotations

import os
import sys
import tempfile

from docx import Document

from core.perfiles import opciones_de_perfil
from core.transformador import adaptar_documento
import crear_ejemplo


def _aprox(valor: float, esperado: float, tol: float = 0.02) -> bool:
    return abs(valor - esperado) <= tol


def main() -> int:
    fallos: list[str] = []
    trabajo = tempfile.mkdtemp(prefix="adaptadocs-prueba-")
    origen = os.path.join(trabajo, "ejemplo.docx")

    cwd = os.getcwd()
    os.chdir(trabajo)
    try:
        crear_ejemplo.main()
    finally:
        os.chdir(cwd)

    # --- Perfil Dislexia --------------------------------------------- #
    salida = os.path.join(trabajo, "dislexia.docx")
    opciones = opciones_de_perfil("Dislexia")
    opciones.resaltar_palabras = ["glucosa", "oxígeno", "estomas"]
    resumen = adaptar_documento(origen, salida, opciones)

    doc = Document(salida)

    if not _aprox(doc.sections[0].left_margin.cm, 3.0):
        fallos.append(f"margen izquierdo = {doc.sections[0].left_margin.cm} cm (esperado ~3.0)")

    cuerpo = doc.paragraphs[1]
    run = cuerpo.runs[0]
    if run.font.name != "Verdana":
        fallos.append(f"fuente = {run.font.name!r} (esperado 'Verdana')")
    if run.font.size is None or not _aprox(run.font.size.pt, 14.0, 0.1):
        fallos.append(f"tamaño = {run.font.size} (esperado 14 pt)")
    if cuerpo.paragraph_format.line_spacing not in (1.5,):
        fallos.append(f"interlineado = {cuerpo.paragraph_format.line_spacing} (esperado 1.5)")

    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if cuerpo.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
        fallos.append("el párrafo sigue justificado")

    if resumen["resaltados"] < 3:
        fallos.append(f"palabras resaltadas = {resumen['resaltados']} (esperado >= 3)")

    resaltadas = [
        r.text
        for p in doc.paragraphs
        for r in p.runs
        if r.font.highlight_color is not None
    ]
    if not any("glucosa" in t.lower() for t in resaltadas):
        fallos.append(f"no se resaltó 'glucosa'; resaltadas: {resaltadas}")

    # --- Sopa de letras / crucigrama: no se le toca la fuente ni el
    # espaciado a una celda de una sola letra (encontrado con un documento
    # real: pasar de Arial Black 11,5 pt a Verdana 14 pt en una celda
    # estrecha desajustaba la cuadrícula del ejercicio). --------------- #
    from docx.shared import Pt as _Pt

    from core.transformador import OpcionesAdaptacion, aplicar_formato

    d_letras = Document()
    tabla = d_letras.add_table(rows=1, cols=2)
    celda_letra = tabla.rows[0].cells[0]
    run_letra = celda_letra.paragraphs[0].add_run("G")
    run_letra.font.name = "Arial Black"
    run_letra.font.size = _Pt(11.5)
    celda_normal = tabla.rows[0].cells[1]
    celda_normal.paragraphs[0].add_run("Instrucciones del ejercicio.")
    ruta_letras = os.path.join(trabajo, "letras.docx")
    d_letras.save(ruta_letras)

    d_letras2 = Document(ruta_letras)
    aplicar_formato(d_letras2, opciones_de_perfil("Dislexia"))
    celda_letra2 = d_letras2.tables[0].rows[0].cells[0]
    run_tras = celda_letra2.paragraphs[0].runs[0]
    if run_tras.font.name != "Arial Black" or not _aprox(run_tras.font.size.pt, 11.5, 0.1):
        fallos.append(
            f"la celda de una sola letra ha cambiado: fuente={run_tras.font.name} "
            f"tamaño={run_tras.font.size}"
        )
    celda_normal2 = d_letras2.tables[0].rows[0].cells[1]
    run_normal2 = celda_normal2.paragraphs[0].runs[0]
    if run_normal2.font.name != "Verdana":
        fallos.append("la celda con una instrucción normal debería sí cambiar de fuente")

    # --- Columnas de ancho desigual (un esquema/mapa conceptual): no se
    # aplanan a una sola columna aunque "una_columna" esté activo. Las de
    # ancho igual (texto normal repartido en columnas) sí se aplanan. ---- #
    from docx.oxml import OxmlElement as _OxmlElement
    from docx.oxml.ns import qn as _qn

    def _poner_columnas(doc_, num, igual):
        sectPr = doc_.sections[0]._sectPr
        existente = sectPr.find(_qn("w:cols"))
        if existente is not None:
            sectPr.remove(existente)
        cols = _OxmlElement("w:cols")
        cols.set(_qn("w:num"), str(num))
        if not igual:
            cols.set(_qn("w:equalWidth"), "0")
        sectPr.append(cols)

    d_esquema = Document()
    d_esquema.add_paragraph("Primario.")
    _poner_columnas(d_esquema, 3, igual=False)
    ruta_esquema = os.path.join(trabajo, "esquema.docx")
    d_esquema.save(ruta_esquema)
    d_esquema2 = Document(ruta_esquema)
    aplicar_formato(d_esquema2, OpcionesAdaptacion(una_columna=True))
    cols_esq = d_esquema2.sections[0]._sectPr.find(_qn("w:cols"))
    if cols_esq is None or cols_esq.get(_qn("w:num")) != "3":
        fallos.append(
            "una sección de columnas de ancho desigual (un esquema) no debería aplanarse: "
            f"num quedó en {cols_esq.get(_qn('w:num')) if cols_esq is not None else None}"
        )

    d_periodico = Document()
    d_periodico.add_paragraph("Texto normal repartido en columnas de periódico.")
    _poner_columnas(d_periodico, 2, igual=True)
    ruta_periodico = os.path.join(trabajo, "periodico.docx")
    d_periodico.save(ruta_periodico)
    d_periodico2 = Document(ruta_periodico)
    aplicar_formato(d_periodico2, OpcionesAdaptacion(una_columna=True))
    cols_per = d_periodico2.sections[0]._sectPr.find(_qn("w:cols"))
    if cols_per is not None and cols_per.get(_qn("w:num"), "1") != "1":
        fallos.append("una sección de columnas de ancho igual (texto normal) debería aplanarse a una")

    # --- Sangría enorme (una franja de la página reservada para una imagen
    # de portada): no se le agranda la fuente, aunque el párrafo tenga
    # varias palabras -no es una sola letra, así que _es_letra_suelta no lo
    # detecta; hace falta la comprobación general de ancho disponible-
    # (encontrado en un documento real: "Comunidad" se partía en sílabas). #
    from docx.shared import Cm as _Cm

    d_sangria = Document()
    p_estrecho = d_sangria.add_paragraph("La economía de la Comunidad de Madrid")
    p_estrecho.paragraph_format.right_indent = _Cm(14.2)
    p_ancho = d_sangria.add_paragraph("Un párrafo normal sin ninguna sangría especial.")
    ruta_sangria = os.path.join(trabajo, "sangria.docx")
    d_sangria.save(ruta_sangria)

    d_sangria2 = Document(ruta_sangria)
    aplicar_formato(d_sangria2, opciones_de_perfil("Dislexia"))
    run_estrecho = d_sangria2.paragraphs[0].runs[0]
    if run_estrecho.font.name == "Verdana":
        fallos.append("un párrafo con sangría enorme (sin sitio real) no debería agrandarse")
    run_ancho = d_sangria2.paragraphs[1].runs[0]
    if run_ancho.font.name != "Verdana":
        fallos.append("un párrafo normal, sin sangría, sí debería agrandarse")

    # --- Aviso de maquetación compleja: un documento normal no debería
    # avisar; uno con muchas secciones, celdas de una sola letra, columnas
    # desiguales y sangrías enormes a la vez sí debería dar las 4 señales. #
    from core.transformador import (
        _UMBRAL_COLUMNAS_ESTRECHAS,
        _UMBRAL_LETRAS_SUELTAS,
        _UMBRAL_SECCIONES,
        detectar_maquetacion_compleja,
    )

    avisos_normal = detectar_maquetacion_compleja(Document(origen), opciones_de_perfil("Dislexia"))
    if avisos_normal:
        fallos.append(f"un documento normal no debería dar avisos de maquetación: {avisos_normal}")

    d_complejo = Document()
    for _ in range(_UMBRAL_SECCIONES + 1):
        d_complejo.add_section()
    tabla_complejo = d_complejo.add_table(rows=1, cols=_UMBRAL_LETRAS_SUELTAS + 1)
    for i in range(_UMBRAL_LETRAS_SUELTAS + 1):
        tabla_complejo.rows[0].cells[i].paragraphs[0].add_run("A")
    for _ in range(_UMBRAL_COLUMNAS_ESTRECHAS + 1):
        p_estrecho_c = d_complejo.add_paragraph("Un párrafo con poco sitio disponible de verdad.")
        p_estrecho_c.paragraph_format.right_indent = _Cm(14.2)
    _poner_columnas(d_complejo, 3, igual=False)
    ruta_complejo = os.path.join(trabajo, "complejo.docx")
    d_complejo.save(ruta_complejo)

    avisos_complejo = detectar_maquetacion_compleja(Document(ruta_complejo), opciones_de_perfil("Dislexia"))
    if len(avisos_complejo) < 4:
        fallos.append(
            f"esperaba las 4 señales de maquetación compleja (secciones, letras sueltas, "
            f"columnas desiguales, columnas estrechas); salieron {len(avisos_complejo)}: {avisos_complejo}"
        )

    # --- Perfil TDAH: viñetas -> lista numerada --------------------- #
    salida_tdah = os.path.join(trabajo, "tdah.docx")
    resumen_tdah = adaptar_documento(origen, salida_tdah, opciones_de_perfil("TDAH"))
    if resumen_tdah["vinetas_convertidas"] < 1:
        fallos.append("no se convirtió ninguna viñeta en lista numerada (perfil TDAH)")

    # --- Separar procedimientos en pasos (sin IA) ------------------- #
    from core.transformador import OpcionesAdaptacion, aplicar_formato

    d_marca = Document(origen)
    r_marca = aplicar_formato(d_marca, OpcionesAdaptacion(separar_en_pasos="marcados"))
    if r_marca["procedimientos_en_pasos"] != 1:
        fallos.append(
            f"marcados: procedimientos_en_pasos = {r_marca['procedimientos_en_pasos']} (esperado 1)"
        )
    numeradas = [p.text for p in d_marca.paragraphs if "Number" in (p.style.name or "")]
    if not any(t.startswith("Coge una planta") for t in numeradas):
        fallos.append(f"marcados: no se troceó el párrafo «PASOS:»; numeradas: {numeradas}")
    if any("PASOS:" in p.text for p in d_marca.paragraphs):
        fallos.append("marcados: no se quitó la marca «PASOS:»")

    d_auto = Document(origen)
    r_auto = aplicar_formato(d_auto, OpcionesAdaptacion(separar_en_pasos="auto"))
    if r_auto["procedimientos_en_pasos"] < 2:
        fallos.append(
            f"auto: procedimientos_en_pasos = {r_auto['procedimientos_en_pasos']} (esperado >= 2)"
        )

    # Cada procedimiento debe estrenar su propia numeración (reinicia en 1),
    # no continuar la del anterior. Se comprueba que los pasos usan al menos
    # dos listas distintas y que cada una arranca con startOverride = 1.
    num_ids_pasos = [
        int(v)
        for p in d_auto.paragraphs
        if "Number" in (p.style.name or "")
        for v in p._p.xpath(".//w:numPr/w:numId/@w:val")
    ]
    if len(set(num_ids_pasos)) < 2:
        fallos.append(
            f"auto: los pasos comparten numeración {set(num_ids_pasos)} "
            "(cada actividad debería reiniciar en 1)"
        )
    numbering = d_auto.part.numbering_part.element
    for nid in set(num_ids_pasos):
        num = numbering.num_having_numId(nid)
        arranques = num.xpath(".//w:lvlOverride/w:startOverride/@w:val")
        if "1" not in arranques:
            fallos.append(f"auto: la lista {nid} no reinicia en 1 (startOverride={arranques})")

    d_no = Document(origen)
    r_no = aplicar_formato(d_no, OpcionesAdaptacion(separar_en_pasos="no"))
    if r_no["procedimientos_en_pasos"] != 0:
        fallos.append("no: no debería haber tocado ningún procedimiento")

    # --- Numerar preguntas y dejar espacio para responder ----------- #
    d_preg_off = Document(origen)
    r_preg_off = aplicar_formato(d_preg_off, OpcionesAdaptacion())
    if r_preg_off["preguntas_numeradas"] != 0 or r_preg_off["preguntas_con_espacio"] != 0:
        fallos.append("por defecto no debería tocarse ninguna pregunta")

    d_preg_num = Document(origen)
    r_preg_num = aplicar_formato(d_preg_num, OpcionesAdaptacion(numerar_preguntas=True))
    if r_preg_num["preguntas_numeradas"] != 2:
        fallos.append(
            f"preguntas_numeradas = {r_preg_num['preguntas_numeradas']} (esperado 2)"
        )
    preguntas_numeradas = [
        p.text for p in d_preg_num.paragraphs if "Number" in (p.style.name or "")
    ]
    if not any(t.startswith("¿Qué necesita una planta") for t in preguntas_numeradas):
        fallos.append("no se numeró la primera pregunta")

    parrafos_antes = len(Document(origen).paragraphs)
    d_preg_esp = Document(origen)
    r_preg_esp = aplicar_formato(d_preg_esp, OpcionesAdaptacion(espacio_respuestas=2))
    if r_preg_esp["preguntas_con_espacio"] != 2:
        fallos.append(
            f"preguntas_con_espacio = {r_preg_esp['preguntas_con_espacio']} (esperado 2)"
        )
    if len(d_preg_esp.paragraphs) != parrafos_antes + 4:
        fallos.append(
            f"no se añadieron las líneas en blanco esperadas "
            f"({len(d_preg_esp.paragraphs)} párrafos, esperado {parrafos_antes + 4})"
        )

    # --- Enunciados de examen (Calcula, Opera, Ordena…) ------------- #
    ex = Document()
    ex.add_paragraph("Nota: justifica todos los pasos.")
    ex.add_paragraph("Calcula el resultado de las operaciones.")
    ex.add_paragraph("- 2 + 2")
    ex.add_paragraph("- 5 x 3")
    ex.add_paragraph("1. Opera y simplifica la expresión.")
    ex.add_paragraph("Ordena de menor a mayor los siguientes números.")
    ruta_ex = os.path.join(trabajo, "examen.docx")
    ex.save(ruta_ex)

    d_ex_off = Document(ruta_ex)
    if aplicar_formato(d_ex_off, OpcionesAdaptacion(espacio_respuestas=3))["preguntas_con_espacio"] != 0:
        fallos.append("sin 'enunciados_examen', los enunciados sin «?» no deberían llevar hueco")

    d_ex = Document(ruta_ex)
    r_ex = aplicar_formato(d_ex, OpcionesAdaptacion(espacio_respuestas=3, enunciados_examen=True))
    if r_ex["preguntas_con_espacio"] != 5:
        fallos.append(
            f"enunciados_examen: preguntas_con_espacio = {r_ex['preguntas_con_espacio']} "
            "(esperado 5: Calcula, 2 guiones, «1. Opera», Ordena)"
        )
    textos_ex = [p.text for p in d_ex.paragraphs]
    if textos_ex.count("") < 15:
        fallos.append(f"enunciados_examen: faltan líneas de respuesta (vacías={textos_ex.count('')})")

    # --- El original no se toca ------------------------------------- #
    if _aprox(Document(origen).sections[0].left_margin.cm, 3.0):
        fallos.append("¡el documento original ha sido modificado!")

    if fallos:
        print("PRUEBA FALLIDA:")
        for f in fallos:
            print("  -", f)
        return 1

    print(
        "PRUEBA OK — formato, resaltado, viñetas, pasos, preguntas, "
        "celdas de una sola letra y columnas desiguales sin tocar, aviso de "
        "maquetación compleja, y original intacto."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
