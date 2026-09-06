"""Comprobación del generador del Anexo III.b (ACIS).

    python prueba_acis.py

Genera el documento con datos de ejemplo y verifica: la estructura oficial,
la casilla legal que deja fuera las materias sin ACS determinada, las marcas
de pendiente, el aviso de borrador y que no se escriben datos del alumno.
Sale con código 0 si todo va bien, 1 si algo falla. No usa conexión.
"""

from __future__ import annotations

import os
import sys
import tempfile

from docx import Document

from core.acis import (
    DatosACIS,
    MateriaACIS,
    cargar_textos,
    datos_desde_dict,
    generar_acis,
    materia_desde_programacion,
)
from core.programacion import Competencia, Programacion


def _texto(doc) -> str:
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for fila in t.rows:
            for celda in fila.cells:
                partes.append(celda.text)
    return "\n".join(partes)


def main() -> int:
    fallos: list[str] = []
    trabajo = tempfile.mkdtemp(prefix="adaptadocs-acis-")

    textos = cargar_textos()
    for clave in ("anexo", "titulo", "marca_borrador", "pendiente", "cuerpo",
                  "nota_expediente", "footer", "aviso_sin_acs"):
        if clave not in textos:
            fallos.append(f"falta la clave «{clave}» en acis.json")

    # --- El diccionario de datos se convierte bien ------------------- #
    datos = datos_desde_dict({
        "materias": [
            {
                "materia": "Biología y Geología",
                "profesor": "Ana P.",
                "departamento": "Biología y Geología",
                "acs_determinada": True,
                "criterios_evaluacion": "1.1. (2.º ESO) Identificar seres vivos.",
                "contenidos": "La célula (2.º ESO).",
                "metodologia": "Apoyo de PT 2 h/semana.",
                "instrumentos": "Pruebas orales; rúbricas.",
                "unidades": "UD 1 a 6 con adaptación.",
                "secuenciacion": [["UD 1. La célula", "1er trimestre"]],
            },
            {"materia": "Matemáticas", "acs_determinada": False},
        ],
    })
    if len(datos.materias) != 2:
        fallos.append(f"se esperaban 2 materias en los datos; hay {len(datos.materias)}")

    # --- Caso normal: una materia con ACS, otra sin ---------------- #
    salida = os.path.join(trabajo, "ACIS.docx")
    ruta, avisos = generar_acis(salida, datos)
    doc = Document(ruta)
    cuerpo = _texto(doc)

    if len(avisos) != 1 or "Matemáticas" not in avisos[0]:
        fallos.append(f"debería haber 1 aviso por «Matemáticas»; avisos={avisos}")
    if avisos and "Decreto 23/2023" not in avisos[0]:
        fallos.append("el aviso no cita la base legal (Decreto 23/2023)")

    # Solo debe generarse el bloque de la materia con ACS (datos alumno + 1 materia)
    if len(doc.tables) != 2:
        fallos.append(f"se esperaban 2 tablas (alumno + 1 materia); hay {len(doc.tables)}")

    comprobaciones = {
        "el anexo": textos["anexo"] in cuerpo,
        "el aviso de borrador": "BORRADOR" in cuerpo,
        "la materia con ACS": "Biología y Geología" in cuerpo,
        "los criterios de evaluación aportados": "Identificar seres vivos" in cuerpo,
        "la secuenciación": "UD 1. La célula" in cuerpo and "1er trimestre" in cuerpo,
        "las competencias vacías como pendiente": textos["pendiente"] in cuerpo,
        "la materia sin ACS en «no incluidas»": "Materias no incluidas" in cuerpo,
        "la nota de expediente": "EXPEDIENTE ACADÉMICO" in cuerpo,
        "el pie de página": textos["footer"] in doc.sections[0].footer.paragraphs[0].text,
    }
    for que, ok in comprobaciones.items():
        if not ok:
            fallos.append(f"el documento no incluye {que}")

    # No lleva bloque de firmas (el documento se copia en Raíces).
    for marca in ("Fdo.:", "V.º B.º", "JEFE/A DE ESTUDIOS", "EL PROFESOR DE"):
        if marca in cuerpo:
            fallos.append(f"el documento no debería llevar «{marca}» (bloque de firmas eliminado)")

    # No se deben escribir datos del alumno: la tabla de datos son solo
    # etiquetas y líneas en blanco.
    tabla_alumno = doc.tables[0]
    texto_alumno = " ".join(c.text for fila in tabla_alumno.rows for c in fila.cells)
    if "Ana" in texto_alumno or any(ch.isdigit() for ch in texto_alumno):
        fallos.append("la tabla de datos del alumno no debe contener datos reales")

    # --- Todas las materias sin ACS: se genera igual, sin bloques -- #
    solo_datos = DatosACIS(materias=[MateriaACIS(materia="Lengua", acs_determinada=False)])
    ruta2, avisos2 = generar_acis(os.path.join(trabajo, "vacio.docx"), solo_datos)
    doc2 = Document(ruta2)
    if len(doc2.tables) != 1:
        fallos.append(f"sin materias con ACS: solo debería estar la tabla de datos; hay {len(doc2.tables)}")
    if len(avisos2) != 1:
        fallos.append(f"sin materias con ACS: debería haber 1 aviso; avisos={avisos2}")

    # --- Materia completa: no deben quedar campos pendientes ------- #
    completa = DatosACIS(materias=[MateriaACIS(
        materia="Física y Química", profesor="J. R.", departamento="FQ",
        acs_determinada=True, competencias="Sin adaptación de competencias.",
        criterios_evaluacion="1.1 (3.º ESO).", contenidos="La materia (3.º ESO).",
        metodologia="Trabajo cooperativo.", instrumentos="Portafolio.",
        unidades="UD 1 a 4.", secuenciacion=[["UD 1", "1er trimestre"]],
    )])
    ruta3, _ = generar_acis(os.path.join(trabajo, "completa.docx"), completa)
    tabla_materia = Document(ruta3).tables[1]
    texto_materia = "\n".join(c.text for fila in tabla_materia.rows for c in fila.cells)
    if textos["pendiente"] in texto_materia:
        fallos.append("una materia con todos los campos rellenos no debería tener pendientes")

    # --- Volcado desde una programación (curso destino) ------------- #
    prog = Programacion(
        materia="Lengua Extranjera (Inglés)", curso="2.º ESO",
        competencias=[
            Competencia("1", "Comprender textos orales y escritos sencillos.",
                        ["CCL2"], [("1.1.", "Captar la información esencial de textos breves y claros.")]),
        ],
        saberes_basicos={"A. Comunicación": ["Funciones: describir, narrar, pedir información."]},
        instrumentos=[("Listening", "Prueba oral"), ("Writing", "Rúbrica")],
    )
    m = materia_desde_programacion(prog, MateriaACIS(materia="Inglés", acs_determinada=True))
    if "Competencia específica 1" not in m.criterios_evaluacion or "1.1." not in m.criterios_evaluacion:
        fallos.append("materia_desde_programacion no vuelca los criterios")
    if "A. Comunicación" not in m.contenidos or "− Funciones" not in m.contenidos:
        fallos.append("materia_desde_programacion no vuelca los contenidos")
    if "Listening: Prueba oral" not in m.instrumentos:
        fallos.append("materia_desde_programacion no vuelca los instrumentos")
    if m.acs_determinada is not True:
        fallos.append("materia_desde_programacion no debe tocar acs_determinada de la base")

    if fallos:
        print("PRUEBA ACIS FALLIDA:")
        for f in fallos:
            print("  -", f)
        return 1

    print("PRUEBA ACIS OK — estructura oficial, casilla legal, pendientes, borrador y sin datos del alumno.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
