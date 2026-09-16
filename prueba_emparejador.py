"""Comprobación (sin red) de `core.emparejador`: la regla acordada es
"misma materia + misma competencia específica + mismo número de criterio",
y si no hay ese número exacto en el nivel de referencia, se marca como sin
equivalente -nunca se inventa nada aquí-. Los saberes básicos se emparejan
por título de bloque, no por letra.

    python prueba_emparejador.py
"""

from __future__ import annotations

import sys

from core.acis import DatosACIS, MateriaACIS, generar_acis
from core.emparejador import (
    avisos,
    cargar_tabla_curriculo,
    emparejar_criterios,
    emparejar_programacion,
    emparejar_saberes,
    texto_criterios,
    texto_saberes,
)
from core.programacion import Competencia, Programacion


def main() -> int:
    fallos: list[str] = []

    # --- Tabla de control pequeña, con el caso real que faltaba emparejar:
    # Matemáticas CE2 tiene DOS criterios en el curso de origen (2.1 y 2.2)
    # pero solo UNO en el nivel de referencia (2.1). ---
    tabla = {
        "criterios": [
            {"materia": "MATEMÁTICAS", "curso": "1.º ESO", "competencia_id": "CE2",
             "criterio_id": "2.1", "criterio_texto": "Conocer y aplicar las herramientas básicas.",
             "referencia": "Decreto 65/2022 — Matemáticas, 1.º ESO"},
            {"materia": "MATEMÁTICAS", "curso": "3.º ESO", "competencia_id": "CE2",
             "criterio_id": "2.1", "criterio_texto": "Comprobar la corrección matemática.",
             "referencia": "Decreto 65/2022 — Matemáticas, 3.º ESO"},
        ],
        "saberes": [
            {"materia": "MATEMÁTICAS", "curso": "1.º ESO", "bloque_id": "B",
             "bloque_titulo": "Medida y geometría", "item": "Perímetros y áreas de figuras planas.",
             "referencia": "Decreto 65/2022 — Matemáticas, 1.º ESO"},
        ],
    }

    origen = [
        Competencia(numero="CE2", texto="Razonar y resolver.",
                    criterios=[("2.1.", "Comprobar la corrección matemática de las soluciones.")]),
    ]
    # Añadimos aparte el criterio 2.2, que en la tabla de control NO existe
    # en 1.º ESO -este es el caso de "sin equivalente".
    origen[0].criterios.append(("2.2.", "Comprobar la validez de las soluciones."))

    resultado = emparejar_criterios("MATEMÁTICAS", origen, "1.º ESO", tabla)
    if len(resultado) != 2:
        fallos.append(f"se esperaban 2 filas de resultado, salieron {len(resultado)}")
    else:
        r21, r22 = resultado
        if not r21.encontrado:
            fallos.append("2.1 debería encontrar equivalente en 1.º ESO")
        elif r21.texto_referencia != "Conocer y aplicar las herramientas básicas.":
            fallos.append(f"2.1: texto de referencia equivocado: «{r21.texto_referencia}»")
        elif "1.º ESO" not in r21.referencia:
            fallos.append(f"2.1: la referencia no cita el curso: «{r21.referencia}»")
        if r22.encontrado:
            fallos.append("2.2 NO debería encontrar equivalente (no existe en 1.º ESO) — no hay que inventarlo")
        elif r22.texto_referencia:
            fallos.append("2.2 sin equivalente no debería llevar texto de referencia")

    # El número de competencia puede venir SIN el prefijo "CE" (según la
    # costumbre del profesor que escribió la programación) y el criterio
    # con el punto final que añade `_trocear_criterios`: debe seguir emparejando igual.
    origen_sin_prefijo = [Competencia(numero="2", texto="...", criterios=[("2.1.", "...")])]
    r = emparejar_criterios("MATEMÁTICAS", origen_sin_prefijo, "1.º ESO", tabla)
    if not r or not r[0].encontrado:
        fallos.append("no empareja cuando el número de competencia viene sin el prefijo «CE»")

    # Materia que no está en la tabla: todo sin equivalente, no debe petar.
    r_otra = emparejar_criterios("MÚSICA", origen, "1.º ESO", tabla)
    if any(x.encontrado for x in r_otra):
        fallos.append("una materia que no está en la tabla no debería encontrar nada")

    # --- Saberes básicos: por título de bloque, no por letra ---
    saberes_origen = {"C. Medida y geometría": ["Área y perímetro de polígonos regulares."]}
    r_sab = emparejar_saberes("MATEMÁTICAS", saberes_origen, "1.º ESO", tabla)
    if len(r_sab) != 1 or not r_sab[0].encontrado:
        fallos.append("el bloque 'Medida y geometría' debería emparejar aunque cambie de letra (C -> B)")
    elif "Perímetros y áreas" not in r_sab[0].items_referencia[0]:
        fallos.append(f"el bloque emparejado no trae los items de referencia esperados: {r_sab[0].items_referencia}")

    saberes_sin_equivalente = {"D. Probabilidad y estadística": ["Media y mediana."]}
    r_sab2 = emparejar_saberes("MATEMÁTICAS", saberes_sin_equivalente, "1.º ESO", tabla)
    if r_sab2[0].encontrado:
        fallos.append("un bloque sin ese título en el nivel de referencia no debería encontrar nada")

    # --- Sanity check contra los datos reales del repositorio: el caso
    # concreto que se discutió con el autor (Matemáticas CE2, 3.º -> 1.º ESO). ---
    tabla_real = cargar_tabla_curriculo()
    origen_real = [Competencia(numero="CE2", texto="...", criterios=[
        ("2.1.", "..."), ("2.2.", "..."),
    ])]
    r_real = emparejar_criterios("MATEMÁTICAS", origen_real, "1.º ESO", tabla_real)
    if not r_real[0].encontrado:
        fallos.append("(datos reales) 2.1 de Matemáticas debería emparejar entre 3.º y 1.º ESO")
    if r_real[1].encontrado:
        fallos.append("(datos reales) 2.2 de Matemáticas NO debería tener equivalente en 1.º ESO")

    # --- Texto para el documento: cita literal, pendiente si no hay
    # equivalente (nunca en blanco), y el hueco de redacción alternativa. ---
    texto_cr = texto_criterios(resultado)
    if "CITA LITERAL" not in texto_cr or "Conocer y aplicar las herramientas básicas" not in texto_cr:
        fallos.append("texto_criterios no muestra el texto literal del criterio encontrado")
    if "SIN EQUIVALENTE" not in texto_cr or "pendiente de generar" not in texto_cr:
        fallos.append("texto_criterios no marca como pendiente el criterio sin equivalente")
    if texto_cr.count("Redacción alternativa del profesor") != 2:
        fallos.append("texto_criterios debería llevar un hueco de redacción alternativa POR CADA criterio")

    texto_sa = texto_saberes(r_sab + r_sab2)
    if "CITA LITERAL" not in texto_sa or "Perímetros y áreas" not in texto_sa:
        fallos.append("texto_saberes no muestra el texto literal del bloque encontrado")
    if "SIN EQUIVALENTE" not in texto_sa:
        fallos.append("texto_saberes no marca como pendiente el bloque sin equivalente")

    lista_avisos = avisos(resultado, r_sab + r_sab2)
    if len(lista_avisos) != 2:  # el criterio 2.2 y el bloque de Probabilidad y estadística
        fallos.append(f"avisos() debería devolver 2 avisos (uno por hueco), salieron {len(lista_avisos)}")
    if not any("2.2" in a for a in lista_avisos) or not any("Probabilidad y estadística" in a for a in lista_avisos):
        fallos.append("los avisos no identifican claramente qué criterio/bloque falta")

    # --- emparejar_programacion(): junta todo en una MateriaACIS real, y el
    # documento final la recoge (incluida la marca de redacción alternativa,
    # que no debe verse como texto suelto -debe quedar resaltada, no impresa
    # tal cual "§REDACCION§..."). ---
    prog = Programacion(
        materia="Matemáticas", curso="3.º ESO",
        competencias=origen,
        saberes_basicos=saberes_origen | saberes_sin_equivalente,
    )
    base = MateriaACIS(materia="Matemáticas", profesor="—", departamento="Matemáticas", acs_determinada=True)
    materia_final = emparejar_programacion("MATEMÁTICAS", prog, "1.º ESO", tabla, base)
    if not materia_final.avisos_ia:
        fallos.append("emparejar_programacion no traslada los avisos a la MateriaACIS")
    if "§REDACCION§" not in materia_final.criterios_evaluacion:
        fallos.append("emparejar_programacion debería dejar la marca interna en el texto (la quita el renderizador)")

    import os
    import tempfile
    from docx import Document
    trabajo = tempfile.mkdtemp(prefix="adaptadocs-emparejador-")
    ruta = os.path.join(trabajo, "ACIS.docx")
    generar_acis(ruta, DatosACIS(materias=[materia_final]))
    doc = Document(ruta)
    texto_documento = "\n".join(p.text for tabla_doc in doc.tables for fila in tabla_doc.rows
                                 for c in fila.cells for p in c.paragraphs)
    if "§REDACCION§" in texto_documento:
        fallos.append("la marca interna «§REDACCION§» no debería verse en el documento final, solo el resaltado")
    if "Redacción alternativa del profesor" not in texto_documento:
        fallos.append("el hueco de redacción alternativa no llega al documento final")
    if "SIN EQUIVALENTE" not in texto_documento:
        fallos.append("el aviso de «sin equivalente» no llega al documento final")

    if fallos:
        print("PRUEBA EMPAREJADOR FALLIDA:")
        for f in fallos:
            print("  -", f)
        return 1

    print("PRUEBA EMPAREJADOR OK — emparejamiento literal, criterios sin equivalente, "
          "prefijo CE opcional, saberes básicos por título, caso real de Matemáticas, "
          "texto del documento, avisos y MateriaACIS completa.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
