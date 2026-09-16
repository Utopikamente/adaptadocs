"""Comprobación (sin red) de `core.emparejador_ia`: construcción del mensaje
a la IA para los huecos, volcado de una respuesta simulada sobre los
criterios/saberes pendientes (sin inventar nada si la IA no propone algo
para uno de ellos), y que `emparejar_programacion_con_ia` no llama a la API
en absoluto cuando no hay ningún hueco. La llamada real se prueba a mano.

    python prueba_emparejador_ia.py
"""

from __future__ import annotations

import sys

from core.acis import DatosACIS, MateriaACIS, generar_acis
from core.emparejador import CriterioEmparejado, SaberEmparejado
from core.emparejador_ia import _mensaje_huecos, aplicar_adaptacion, emparejar_programacion_con_ia
from core.programacion import Competencia, Programacion


def main() -> int:
    fallos: list[str] = []

    ancla_cr = CriterioEmparejado(
        numero="2.1.", competencia="CE2", texto_origen="...", encontrado=True,
        texto_referencia="Conocer y aplicar las herramientas básicas.",
        referencia="Decreto 65/2022 — Matemáticas, 1.º ESO",
    )
    ancla_sa = SaberEmparejado(
        bloque_origen="B. Medida y geometría", titulo="Medida y geometría",
        items_origen=[], encontrado=True, items_referencia=["Perímetros y áreas de figuras planas."],
    )
    pendiente_cr = CriterioEmparejado(
        numero="2.2.", competencia="CE2",
        texto_origen="Comprobar la validez de las soluciones de un problema.", encontrado=False,
    )
    pendiente_sa = SaberEmparejado(
        bloque_origen="C. Geometría en el plano y el espacio", titulo="Geometría en el plano y el espacio",
        items_origen=["Teorema de Pitágoras."], encontrado=False,
    )

    # --- El mensaje a la IA debe llevar las anclas (como modelo) y los
    # huecos (a adaptar), pero nunca mezclarlos ni pedir que se reescriban
    # las anclas. ---
    msg = _mensaje_huecos("MATEMÁTICAS", "3.º ESO", "1.º ESO",
                          [pendiente_cr], [pendiente_sa], [ancla_cr], [ancla_sa])
    if "Conocer y aplicar las herramientas básicas" not in msg:
        fallos.append("el mensaje no incluye el criterio ancla (literal) como modelo")
    if "Comprobar la validez de las soluciones" not in msg:
        fallos.append("el mensaje no incluye el criterio pendiente a adaptar")
    if "Perímetros y áreas" not in msg or "Teorema de Pitágoras" not in msg:
        fallos.append("el mensaje no incluye ambos saberes básicos (ancla y pendiente)")
    if "3.º ESO" not in msg or "1.º ESO" not in msg:
        fallos.append("el mensaje no deja claro el curso de origen y el de referencia")

    # --- Volcado de una respuesta simulada de la IA sobre los pendientes. ---
    resultado_simulado = {
        "criterios_adaptados": [
            {"numero": "2.2.", "texto": "Comprobar si el resultado de un problema tiene sentido."},
        ],
        "saberes_adaptados": [
            {"titulo": "Geometría en el plano y el espacio", "items": ["Formas geométricas básicas del entorno."]},
        ],
    }
    aplicar_adaptacion("3.º ESO", "1.º ESO", [pendiente_cr], [pendiente_sa], resultado_simulado)

    if not pendiente_cr.adaptado or pendiente_cr.texto_referencia != "Comprobar si el resultado de un problema tiene sentido.":
        fallos.append("aplicar_adaptacion no ha volcado el texto adaptado del criterio")
    if "Propuesta adaptada" not in pendiente_cr.referencia or "3.º ESO" not in pendiente_cr.referencia:
        fallos.append(f"la referencia del criterio adaptado no es clara: «{pendiente_cr.referencia}»")
    if not pendiente_sa.adaptado or pendiente_sa.items_referencia != ["Formas geométricas básicas del entorno."]:
        fallos.append("aplicar_adaptacion no ha volcado los items adaptados del bloque de saberes")

    # Si la IA no propone nada para un elemento, se queda pendiente -no se inventa nada-.
    otro_pendiente = CriterioEmparejado(numero="9.9.", competencia="CE9", texto_origen="...", encontrado=False)
    aplicar_adaptacion("3.º ESO", "1.º ESO", [otro_pendiente], [], {"criterios_adaptados": [], "saberes_adaptados": []})
    if otro_pendiente.adaptado or otro_pendiente.texto_referencia:
        fallos.append("si la IA no propone nada para un criterio, no debería quedar marcado como adaptado")

    # --- El documento final distingue lo adaptado de lo literal y de lo pendiente. ---
    base = MateriaACIS(materia="Matemáticas", profesor="—", departamento="Matemáticas", acs_determinada=True)
    from core.emparejador import texto_criterios, texto_saberes, avisos as avisos_fn
    base.criterios_evaluacion = texto_criterios([ancla_cr, pendiente_cr])
    base.contenidos = texto_saberes([ancla_sa, pendiente_sa])
    base.avisos_ia = avisos_fn([ancla_cr, pendiente_cr], [ancla_sa, pendiente_sa])

    import os
    import tempfile
    from docx import Document
    trabajo = tempfile.mkdtemp(prefix="adaptadocs-emparejador-ia-")
    ruta = os.path.join(trabajo, "ACIS.docx")
    generar_acis(ruta, DatosACIS(materias=[base]))
    texto_doc = "\n".join(p.text for t in Document(ruta).tables for fila in t.rows
                           for c in fila.cells for p in c.paragraphs)
    if "CITA LITERAL" not in texto_doc:
        fallos.append("el documento no distingue la cita literal")
    if "Comprobar si el resultado de un problema tiene sentido" not in texto_doc:
        fallos.append("el criterio adaptado no llega al documento final")
    if "Propuesta adaptada a partir del criterio 2.2." not in texto_doc:
        fallos.append("el documento no deja claro que el 2.2 es una propuesta adaptada, no literal")
    if "se ha generado una propuesta adaptada" not in texto_doc:
        fallos.append("el aviso amarillo no distingue lo adaptado de lo todavía pendiente")

    # --- Sin ningún hueco, emparejar_programacion_con_ia no debe necesitar
    # (ni intentar usar) la API en absoluto. ---
    tabla = {
        "criterios": [
            {"materia": "MATEMÁTICAS", "curso": "1.º ESO", "competencia_id": "CE1",
             "criterio_id": "1.1", "criterio_texto": "Un criterio literal.",
             "referencia": "Decreto 65/2022 — Matemáticas, 1.º ESO"},
        ],
        "saberes": [],
    }
    prog = Programacion(
        materia="Matemáticas", curso="3.º ESO",
        competencias=[Competencia(numero="CE1", texto="...", criterios=[("1.1.", "...")])],
    )
    try:
        materia_sin_ia = emparejar_programacion_con_ia("MATEMÁTICAS", prog, "1.º ESO", tabla)
    except Exception as exc:  # noqa: BLE001
        fallos.append(f"sin huecos no debería hacer falta clave de API, pero ha fallado: {exc}")
    else:
        if "Un criterio literal." not in materia_sin_ia.criterios_evaluacion:
            fallos.append("el emparejamiento sin huecos no ha producido el texto esperado")
        if materia_sin_ia.avisos_ia:
            fallos.append("sin huecos no debería haber ningún aviso")

    if fallos:
        print("PRUEBA EMPAREJADOR IA (sin red) FALLIDA:")
        for f in fallos:
            print("  -", f)
        return 1

    print("PRUEBA EMPAREJADOR IA (sin red) OK — mensaje a la IA, volcado de la respuesta, "
          "documento final distingue literal/adaptado/pendiente, y coste cero sin huecos.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
