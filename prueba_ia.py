"""Comprobación de la capa de IA que NO usa la red.

Verifica que `core.aplicar_ia.aplicar_resultado` vuelca correctamente sobre el
documento una respuesta simulada (párrafos reescritos, pasos, glosario, resumen
y preguntas). La llamada real a la API se prueba a mano con `prueba_ia_api.py`.

    python prueba_ia.py
"""

from __future__ import annotations

import sys
import tempfile

from docx import Document

from core.aplicar_ia import aplicar_resultado
from core.ia import OpcionesIA, _mensaje_usuario, es_cabecera_datos_alumno
from core.pipeline import _extraer_bloques
from core.transformador import detectar_niveles_titulo
import crear_ejemplo


def main() -> int:
    fallos: list[str] = []

    trabajo = tempfile.mkdtemp(prefix="adaptadocs-ia-")
    import os

    cwd = os.getcwd()
    os.chdir(trabajo)
    try:
        crear_ejemplo.main()
    finally:
        os.chdir(cwd)

    doc = Document(f"{trabajo}/ejemplo.docx")
    bloques, por_id = _extraer_bloques(doc)

    # Localiza ids reales del documento de ejemplo
    id_parrafo_intro = next(b.id for b in bloques if b.tipo == "parrafo")
    id_titulo = next(b.id for b in bloques if b.tipo == "titulo")

    datos = {
        "parrafos_simplificados": [
            {"id": id_parrafo_intro, "texto": "Las plantas fabrican su comida. Usan el sol."}
        ],
        "parrafos_en_pasos": [],
        "glosario": [
            {"termino": "glucosa", "definicion": "El alimento que fabrica la planta."},
            {"termino": "oxígeno", "definicion": "Un gas que necesitamos para respirar."},
        ],
        "resumen": [
            {"apartado": "La fotosíntesis", "puntos": ["Las plantas hacen su comida.", "Sueltan oxígeno."]}
        ],
        "preguntas": [
            "¿Qué usan las plantas para fabricar su alimento?",
            "¿Qué gas liberan las plantas?",
        ],
    }

    resumen = aplicar_resultado(doc, datos, por_id)

    salida = f"{trabajo}/ejemplo-ia.docx"
    doc.save(salida)
    releido = Document(salida)
    textos = [p.text for p in releido.paragraphs]

    if resumen["simplificados"] != 1:
        fallos.append(f"simplificados = {resumen['simplificados']} (esperado 1)")
    if not any("fabrican su comida" in t for t in textos):
        fallos.append("no se ve el párrafo reescrito en el documento")
    if resumen["glosario"] != 2 or "Glosario" not in textos:
        fallos.append("falta la sección Glosario")
    if "Preguntas de comprensión" not in textos:
        fallos.append("falta la sección de preguntas")
    if resumen["preguntas"] != 2:
        fallos.append(f"preguntas = {resumen['preguntas']} (esperado 2)")
    if "Resumen" not in textos:
        fallos.append("falta la sección Resumen")
    # el resumen debe ir cerca del principio, no al final
    if "Resumen" in textos and textos.index("Resumen") > textos.index("Glosario"):
        fallos.append("el resumen no está al principio")

    # Prueba de 'pasos': un párrafo se convierte en varias entradas de lista
    doc2 = Document(f"{trabajo}/ejemplo.docx")
    bloques2, por_id2 = _extraer_bloques(doc2)
    id_algun_parrafo = next(b.id for b in bloques2 if b.tipo == "parrafo")
    n_parrafos_antes = len(doc2.paragraphs)
    r2 = aplicar_resultado(
        doc2,
        {
            "parrafos_simplificados": [],
            "parrafos_en_pasos": [
                {"id": id_algun_parrafo, "pasos": ["Paso uno.", "Paso dos.", "Paso tres."]}
            ],
            "glosario": [],
            "resumen": [],
            "preguntas": [],
        },
        por_id2,
    )
    if r2["en_pasos"] != 1:
        fallos.append(f"en_pasos = {r2['en_pasos']} (esperado 1)")
    if len(doc2.paragraphs) != n_parrafos_antes + 2:
        fallos.append(
            f"pasos: {len(doc2.paragraphs)} párrafos (esperado {n_parrafos_antes + 2})"
        )

    # Prueba de 'preguntas_divididas': una pregunta compuesta se parte en varias
    doc3 = Document(f"{trabajo}/ejemplo.docx")
    bloques3, por_id3 = _extraer_bloques(doc3)
    id_pregunta = next(
        b.id for b in bloques3 if b.tipo == "parrafo" and b.texto.strip().endswith("?")
    )
    n_parrafos_antes3 = len(doc3.paragraphs)
    r3 = aplicar_resultado(
        doc3,
        {
            "parrafos_simplificados": [],
            "parrafos_en_pasos": [],
            "glosario": [],
            "resumen": [],
            "preguntas": [],
            "preguntas_divididas": [
                {
                    "id": id_pregunta,
                    "subpreguntas": [
                        "¿Qué usa la planta del suelo?",
                        "¿Qué usa la planta del aire?",
                    ],
                }
            ],
        },
        por_id3,
    )
    if r3["preguntas_divididas"] != 1:
        fallos.append(f"preguntas_divididas = {r3['preguntas_divididas']} (esperado 1)")
    if len(doc3.paragraphs) != n_parrafos_antes3 + 1:
        fallos.append(
            f"preguntas_divididas: {len(doc3.paragraphs)} párrafos "
            f"(esperado {n_parrafos_antes3 + 1})"
        )
    textos3 = [p.text for p in doc3.paragraphs]
    if "¿Qué usa la planta del suelo?" not in textos3 or "¿Qué usa la planta del aire?" not in textos3:
        fallos.append("no se ven las dos subpreguntas en el documento")

    # Niveles de título elegibles para el resumen: el docente puede decidir
    # que un nivel de encabezado (p. ej. los "Heading 2" de cada ejercicio)
    # no genere su propio mini-resumen, sin perder la protección de no
    # simplificarlo. `crear_ejemplo.docx` tiene 1 "Heading 1" y 4 "Heading 2".
    doc4 = Document(f"{trabajo}/ejemplo.docx")
    niveles = detectar_niveles_titulo(doc4)
    niveles_por_estilo = {n["estilo"]: n for n in niveles}
    if "Heading 1" not in niveles_por_estilo or "Heading 2" not in niveles_por_estilo:
        fallos.append(f"detectar_niveles_titulo no encuentra los dos niveles esperados: {niveles}")
    elif niveles_por_estilo["Heading 2"]["veces"] != 4:
        fallos.append(f"Heading 2 debería aparecer 4 veces, salió {niveles_por_estilo['Heading 2']['veces']}")
    elif niveles_por_estilo["Heading 1"]["ejemplo"] != "La fotosíntesis":
        fallos.append(f"el ejemplo de Heading 1 no es el esperado: «{niveles_por_estilo['Heading 1']['ejemplo']}»")

    bloques4, _ = _extraer_bloques(doc4, niveles_resumen={"Heading 1"})
    titulos4 = {b.texto: b.resumen_candidato for b in bloques4 if b.tipo == "titulo"}
    if titulos4.get("La fotosíntesis") is not True:
        fallos.append("el título de nivel elegido (Heading 1) debería ser candidato a resumen")
    if titulos4.get("Vocabulario importante") is not False:
        fallos.append("el título de nivel NO elegido (Heading 2) no debería ser candidato a resumen")

    msg = _mensaje_usuario(bloques4, OpcionesIA(resumen=True))
    if "(titulo_secundario) Vocabulario importante" not in msg:
        fallos.append("el mensaje a la IA no etiqueta como titulo_secundario el nivel no elegido")
    if "(titulo) La fotosíntesis" not in msg:
        fallos.append("el mensaje a la IA no mantiene (titulo) en el nivel elegido")

    # Sin elegir niveles (comportamiento de siempre): todos los títulos son candidatos.
    bloques4b, _ = _extraer_bloques(doc4)
    if not all(b.resumen_candidato for b in bloques4b if b.tipo == "titulo"):
        fallos.append("sin elegir niveles, todos los títulos deberían seguir siendo candidatos (compatibilidad)")

    # Cabeceras con datos del alumno: nunca deben mandarse a la IA
    positivos = [
        "Nombre y apellidos: ___________________ Curso: ____ Fecha: ___",
        "Nombre y apellidos: Juan Pérez García     Curso: 3º ESO A",
        "Alumno/a: María López",
        "NIA: 12345678",
    ]
    negativos = [
        "Nombre y clasificación de los reinos de la naturaleza.",
        "En este curso aprenderemos sobre las fracciones y los decimales.",
        "¿Cuál es el nombre del proceso por el que las plantas fabrican su alimento?",
    ]
    for texto in positivos:
        if not es_cabecera_datos_alumno(texto):
            fallos.append(f"debería detectarse como cabecera de datos del alumno: «{texto}»")
    for texto in negativos:
        if es_cabecera_datos_alumno(texto):
            fallos.append(f"NO debería detectarse como cabecera de datos del alumno: «{texto}»")

    doc4 = Document(f"{trabajo}/ejemplo.docx")
    doc4.paragraphs[0].insert_paragraph_before(
        "Nombre y apellidos: ___________________ Curso: ____ Fecha: ___"
    )
    bloques4, _ = _extraer_bloques(doc4)
    if any(es_cabecera_datos_alumno(b.texto) for b in bloques4):
        fallos.append("la cabecera de datos del alumno se ha incluido entre los bloques enviados a la IA")

    if fallos:
        print("PRUEBA IA (sin red) FALLIDA:")
        for f in fallos:
            print("  -", f)
        return 1

    print(
        "PRUEBA IA (sin red) OK — reescritura, pasos, glosario, resumen, preguntas, "
        "preguntas divididas, exclusión de cabeceras con datos del alumno y "
        "selección de niveles de título candidatos a resumen."
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
