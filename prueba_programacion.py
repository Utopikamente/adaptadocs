"""Comprobación del lector de programaciones didácticas.

    python prueba_programacion.py

Construye una programación de ejemplo con la estructura de tablas habitual
(competencias | descriptores | criterios de evaluación, fila de saberes
básicos, y tabla de instrumentos) y verifica que `leer_programacion` la
interpreta bien. Sale con 0 si todo va bien, 1 si algo falla. Sin conexión.
"""

from __future__ import annotations

import os
import sys
import tempfile

from docx import Document

from core.programacion import ProgramacionNoReconocida, leer_programacion


def _crear_programacion_ejemplo(ruta: str) -> None:
    doc = Document()
    doc.add_paragraph("PROGRAMACIÓN DIDÁCTICA — MATERIA DE EJEMPLO, 4.º ESO")
    doc.add_paragraph("4. COMPETENCIAS ESPECÍFICAS DE LA MATERIA")

    tabla = doc.add_table(rows=1, cols=3)
    hdr = tabla.rows[0].cells
    hdr[0].text = "COMPETENCIAS ESPECÍFICAS"
    hdr[1].text = "DESCRIPTORES"
    hdr[2].text = "CRITERIOS DE EVALUACIÓN"

    filas = [
        ("1. Comprender textos orales y escritos sencillos para responder a necesidades comunicativas.",
         "CCL2 / CCL3 / CP1",
         "1.1. Extraer el sentido global de textos orales y escritos claros.\n"
         "1.2. Interpretar el contenido de textos progresivamente más complejos."),
        ("2. Producir textos escritos con una organización clara.",
         "CCL1 / CP2 / CD2",
         "2.1. Redactar textos breves y coherentes sobre asuntos cotidianos."),
    ]
    for c0, c1, c2 in filas:
        celdas = tabla.add_row().cells
        celdas[0].text, celdas[1].text, celdas[2].text = c0, c1, c2

    saberes_hdr = tabla.add_row().cells
    for celda in saberes_hdr:
        celda.text = "SABERES BÁSICOS"
    saberes = tabla.add_row().cells
    texto_saberes = (
        "A. Comunicación\n"
        "Estrategias de comprensión y producción de textos.\n"
        "Funciones comunicativas de uso común: saludar, describir, narrar.\n"
        "B. Plurilingüismo\n"
        "Estrategias para transferir conocimientos entre lenguas."
    )
    for celda in saberes:
        celda.text = texto_saberes

    doc.add_paragraph("7. INSTRUMENTOS DE EVALUACIÓN")
    ti = doc.add_table(rows=1, cols=2)
    ti.rows[0].cells[0].text = "ELEMENTO"
    ti.rows[0].cells[1].text = "INSTRUMENTO DE EVALUACIÓN"
    for elemento, instrumento in [("Listening", "Prueba escrita"), ("Writing", "Rúbrica")]:
        c = ti.add_row().cells
        c[0].text, c[1].text = elemento, instrumento

    doc.save(ruta)


def main() -> int:
    fallos: list[str] = []
    trabajo = tempfile.mkdtemp(prefix="adaptadocs-prog-")
    ruta = os.path.join(trabajo, "programacion.docx")
    _crear_programacion_ejemplo(ruta)

    prog = leer_programacion(ruta, materia="Materia de ejemplo", curso="4.º ESO")

    if len(prog.competencias) != 2:
        fallos.append(f"competencias detectadas = {len(prog.competencias)} (esperado 2)")

    if prog.competencias:
        c1 = prog.competencias[0]
        if c1.numero != "1":
            fallos.append(f"número de la 1.ª competencia = {c1.numero!r} (esperado '1')")
        if "Comprender textos" not in c1.texto:
            fallos.append("no se extrajo el texto de la 1.ª competencia")
        if c1.descriptores != ["CCL2", "CCL3", "CP1"]:
            fallos.append(f"descriptores de la 1.ª competencia = {c1.descriptores}")
        codigos = [cod for cod, _ in c1.criterios]
        if codigos != ["1.1.", "1.2."]:
            fallos.append(f"criterios de la 1.ª competencia = {codigos} (esperado ['1.1.', '1.2.'])")
        if c1.criterios and "sentido global" not in c1.criterios[0][1]:
            fallos.append("no se extrajo el texto del criterio 1.1.")

    bloques = list(prog.saberes_basicos)
    if bloques != ["A. Comunicación", "B. Plurilingüismo"]:
        fallos.append(f"bloques de saberes básicos = {bloques}")
    if prog.saberes_basicos.get("A. Comunicación") and \
            len(prog.saberes_basicos["A. Comunicación"]) != 2:
        fallos.append("el bloque A debería tener 2 ítems")

    if prog.instrumentos != [("Listening", "Prueba escrita"), ("Writing", "Rúbrica")]:
        fallos.append(f"instrumentos = {prog.instrumentos}")

    # --- Programación en otro formato (títulos + prosa, sin la tabla) --- #
    ruta_prosa = os.path.join(trabajo, "prosa.docx")
    d = Document()
    d.add_paragraph("4. COMPETENCIAS ESPECÍFICAS")
    d.add_paragraph("1. Comprender textos orales y escritos sencillos para responder a necesidades.")
    d.add_paragraph("2. Producir textos escritos con una organización clara.")
    d.add_paragraph("5. CRITERIOS DE EVALUACIÓN")
    d.add_paragraph("1.1. Extraer el sentido global de textos claros.")
    d.add_paragraph("1.2. Interpretar textos progresivamente más complejos.")
    d.add_paragraph("2.1. Redactar textos breves y coherentes.")
    d.add_paragraph("6. CONTENIDOS")
    d.add_paragraph("A. Comunicación")
    d.add_paragraph("- Estrategias de comprensión y producción.")
    d.add_paragraph("- Funciones: saludar, describir, narrar.")
    d.add_paragraph("7. INSTRUMENTOS DE EVALUACIÓN Y CRITERIOS DE CALIFICACIÓN")
    d.add_paragraph("Pruebas escritas 60 %, trabajos 30 %, actitud 10 %.")
    d.add_paragraph("8. METODOLOGÍA")
    d.add_paragraph("Enfoque comunicativo, trabajo por tareas y agrupamientos flexibles.")
    d.add_paragraph("9. UNIDADES DIDÁCTICAS")
    d.add_paragraph("UNIT 1 FAMILY MATTERS 1ª evaluación")
    d.add_paragraph("UNIT 2 SWEET DREAMS 1ª evaluación")
    d.add_paragraph("Tema 3. La ciudad — 2º trimestre")
    d.save(ruta_prosa)
    p2 = leer_programacion(ruta_prosa)
    if len(p2.competencias) != 2:
        fallos.append(f"prosa: competencias = {len(p2.competencias)} (esperado 2)")
    if [c for c, _ in p2.criterios] != ["1.1.", "1.2.", "2.1."]:
        fallos.append(f"prosa: criterios = {[c for c, _ in p2.criterios]}")
    if "A. Comunicación" not in p2.saberes_basicos:
        fallos.append("prosa: no detecta el bloque de contenidos «A. Comunicación»")
    if "60 %" not in p2.instrumentos_texto:
        fallos.append("prosa: no recoge los instrumentos / criterios de calificación")
    if "agrupamientos flexibles" not in p2.metodologia:
        fallos.append("prosa: no recoge la metodología")
    titulos = [t for t, _ in p2.unidades]
    if len(p2.unidades) != 3:
        fallos.append(f"prosa: unidades = {len(p2.unidades)} (esperado 3): {titulos}")
    if p2.unidades and p2.unidades[0] != ("UNIT 1 FAMILY MATTERS", "1ª evaluación"):
        fallos.append(f"prosa: primera unidad mal extraída: {p2.unidades[0]}")
    if not any("La ciudad" in t for t in titulos):
        fallos.append("prosa: no reconoce «Tema 3. …» como unidad")

    # --- Documento sin ningún apartado curricular ------------------ #
    ruta_mala = os.path.join(trabajo, "sin_apartados.docx")
    d = Document()
    d.add_paragraph("Un documento cualquiera sin apartados de programación.")
    d.save(ruta_mala)
    try:
        leer_programacion(ruta_mala)
        fallos.append("un documento sin apartados debería lanzar ProgramacionNoReconocida")
    except ProgramacionNoReconocida:
        pass

    if fallos:
        print("PRUEBA PROGRAMACIÓN FALLIDA:")
        for f in fallos:
            print("  -", f)
        return 1

    print("PRUEBA PROGRAMACIÓN OK — competencias, descriptores, criterios, saberes básicos e instrumentos.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
