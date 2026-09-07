"""Comprobación del registro de adaptaciones (documento interno).

    python prueba_registro.py

Genera el documento de registro con datos de ejemplo y verifica que recoge
las adaptaciones aplicadas, el marco legal y el aviso de privacidad, y que no
inventa datos del alumno. Sale con código 0 si todo va bien, 1 si algo falla.
No usa conexión.
"""

from __future__ import annotations

import os
import sys
import tempfile

from docx import Document

from core.registro import cargar_textos, generar_registro, ruta_registro
from core.transformador import OpcionesAdaptacion


def _texto(doc) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def main() -> int:
    fallos: list[str] = []
    trabajo = tempfile.mkdtemp(prefix="adaptadocs-registro-")

    # --- Los textos editables se cargan y tienen las claves esperadas --- #
    textos = cargar_textos()
    for clave in ("titulo", "marco_legal", "nota_pie", "lineas_formato", "lineas_ia"):
        if clave not in textos:
            fallos.append(f"falta la clave «{clave}» en registro.json")

    # --- Nombre del archivo de registro -------------------------------- #
    r = ruta_registro(os.path.join(trabajo, "ficha (adaptado).docx"))
    if os.path.basename(r) != "ficha (adaptado) — registro.docx":
        fallos.append(f"nombre del registro inesperado: {os.path.basename(r)}")

    # --- Caso con adaptaciones de formato ------------------------------ #
    opciones = OpcionesAdaptacion(
        fuente="Verdana", tamano_pt=14, interlineado=1.5, espacio_despues_pt=12,
        alinear_izquierda=True, margenes_cm=3.0, una_columna=True,
        alto_contraste=True, negrita_titulos=True,
    )
    resumen = {
        "resaltados": 7, "vinetas_convertidas": 0, "procedimientos_en_pasos": 2,
        "preguntas_numeradas": 3, "preguntas_con_espacio": 3,
    }
    salida = os.path.join(trabajo, "ficha (adaptado) — registro.docx")
    generar_registro(salida, "ficha (adaptado).docx", opciones, resumen)
    doc = Document(salida)
    cuerpo = _texto(doc)

    comprobaciones = {
        "el título": textos["titulo"] in cuerpo,
        "el nombre del documento adaptado": "ficha (adaptado).docx" in cuerpo,
        "el marco legal (Decreto 23/2023)": "Decreto 23/2023" in cuerpo,
        "el marco legal (Orden 1712/2023)": "Orden 1712/2023" in cuerpo,
        "la tipografía": "Verdana" in cuerpo and "14" in cuerpo,
        "el interlineado 1,5": "1,5" in cuerpo,
        "los márgenes 3 cm": "3 cm" in cuerpo,
        "las palabras resaltadas (7)": "7" in cuerpo,
        "los procedimientos en pasos (2)": "pasos numerados (2)" in cuerpo,
        "las preguntas numeradas (3)": "consecutiva (3)" in cuerpo,
        "el hueco para el nombre del alumno": "Alumno/a: _____" in cuerpo,
        "el aviso de privacidad": "No contiene datos de salud" in cuerpo,
    }
    for que, ok in comprobaciones.items():
        if not ok:
            fallos.append(f"el registro no incluye {que}")

    if "TDAH" in cuerpo or "dislexia" in cuerpo.lower():
        fallos.append("el registro no debe nombrar perfiles ni condiciones clínicas")

    vinetas = [p.text for p in doc.paragraphs if "Bullet" in (p.style.name or "")]
    if len(vinetas) < 5:
        fallos.append(f"se esperaban varias líneas de adaptación; hay {len(vinetas)}")

    # --- Con capa de IA ---------------------------------------------- #
    class _OpcionesIA:
        simplificar = True
        pasos = False
        glosario = True
        resumen = False
        preguntas = True
        dividir_preguntas = True
        nivel = "3º-4º de Primaria"

    salida_ia = os.path.join(trabajo, "ficha ia — registro.docx")
    generar_registro(
        salida_ia, "ficha (adaptado).docx", opciones, resumen,
        resumen_ia={"preguntas": 5, "preguntas_divididas": 4}, opciones_ia=_OpcionesIA(),
    )
    cuerpo_ia = _texto(Document(salida_ia))
    if "3º-4º de Primaria" not in cuerpo_ia:
        fallos.append("el registro con IA no recoge el nivel de lectura")
    if "(4)" not in cuerpo_ia:
        fallos.append("el registro con IA no recoge las preguntas divididas (4)")
    if textos["titulo_ia"] not in cuerpo_ia:
        fallos.append("falta el apartado de contenido con IA")

    # --- Caso sin ninguna adaptación de formato --------------------- #
    vacio = OpcionesAdaptacion(
        fuente=None, tamano_pt=None, interlineado=None, espacio_despues_pt=None,
        alinear_izquierda=False, margenes_cm=None, una_columna=False,
        alto_contraste=False, negrita_titulos=False,
    )
    salida_vacia = os.path.join(trabajo, "vacio — registro.docx")
    generar_registro(salida_vacia, "x.docx", vacio,
                     {"resaltados": 0, "procedimientos_en_pasos": 0,
                      "preguntas_numeradas": 0, "preguntas_con_espacio": 0})
    if textos["sin_adaptaciones"] not in _texto(Document(salida_vacia)):
        fallos.append("sin adaptaciones: debería decirlo explícitamente")

    if fallos:
        print("PRUEBA REGISTRO FALLIDA:")
        for f in fallos:
            print("  -", f)
        return 1

    print("PRUEBA REGISTRO OK — adaptaciones, marco legal, IA, privacidad y sin datos del alumno.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
